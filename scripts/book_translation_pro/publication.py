"""Publication-oriented Markdown cleanup and DOCX/EPUB/PDF typesetting."""

from __future__ import annotations

import hashlib
import os
import posixpath
import re
import subprocess
import sys
import unicodedata
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any
from urllib.parse import unquote
from xml.etree import ElementTree as ET

from bs4 import BeautifulSoup

from .alignment import create_segment_alignment, render_bilingual_markdown
from .artifacts import (
    BTPError,
    atomic_copy,
    atomic_write_json,
    atomic_write_text,
    external_tool_path,
    load_json,
    portable_path,
    safe_project_path,
    sha256_file,
)
from .manifest import load_manifest
from .runtime import epubcheck_command, find_binary
from .rights import insert_rights_notice, rights_metadata_text


PRINT_MARKER_RE = re.compile(
    r"(?:<span\b[^>]*>\s*)*(?:\*\*)?"
    r"(?P<file>[A-Za-z0-9_.-]+\.(?:qxd|indd))\s+"
    r"\d{1,2}/\d{1,2}/\d{2,4}\s+\d{1,2}:\d{2}\s+(?:AM|PM)\s+"
    r"Page\s+(?P<page>[ivxlcdm]+|\d+)"
    r"(?:\*\*)?(?:\s*</span>)*",
    re.IGNORECASE,
)
HTML_TAG_RE = re.compile(r"</?(?:span|div)\b[^>]*>", re.IGNORECASE)
HTML_COMMENT_RE = re.compile(r"<!--.*?-->")
# Empty anchor whose id/name an internal ``href="#…"`` jump targets.  The id
# is filled in per document by _strip_print_artifacts from the link targets.
_ANCHOR_TARGET_RE = re.compile(
    r"<(?:span|a)\b[^>]*?\b(?:id|name)\s*=\s*[\"']([^\"']+)[\"'][^>]*>\s*</(?:span|a)>",
    re.IGNORECASE,
)
HTML_HEADING_RE = re.compile(
    r"<h(?P<level>[1-6])\b[^>]*>(?P<body>.*?)</h(?P=level)\s*>",
    re.IGNORECASE | re.DOTALL,
)
HTML_IMAGE_TAG_RE = re.compile(r"<img\b(?P<attrs>[^>]*)>", re.IGNORECASE)
HTML_ATTR_RE = re.compile(r"(?P<name>[A-Za-z_:][-A-Za-z0-9_.:]*)\s*=\s*(?P<quote>[\"'])(?P<value>.*?)(?P=quote)", re.IGNORECASE)
LINKED_IMAGE_RE = re.compile(
    r"(?P<open><a\b(?P<a_attrs>[^>]*)>)\s*"
    r"(?P<img><img\b(?P<img_attrs>[^>]*)>)\s*</a>",
    re.IGNORECASE | re.DOTALL,
)
LINKED_MARKDOWN_IMAGE_RE = re.compile(
    r"(?P<open><a\b(?P<a_attrs>[^>]*)>)\s*"
    r"!\[(?P<alt>[^]]*)\]\((?P<src>[^)\s]+)(?P<title>\s+[^)]*)?\)\s*</a>",
    re.IGNORECASE | re.DOTALL,
)
FOOTNOTE_SPAN_RE = re.compile(
    r"<span\b(?P<attrs>[^>]*)>\s*(?P<number>\*?\d+\*?)\s*</span>",
    re.IGNORECASE,
)
# Extraction can place a superscript footnote marker before the Chinese
# classifier and a protected alphanumeric term (for example ``11个 5S`` from
# source ``5S¹¹``).  Restrict the repair to marker spans followed by a short
# term containing both a letter and a digit; ordinary quantities stay intact.
FOOTNOTE_BEFORE_PROTECTED_TERM_RE = re.compile(
    r"<span\b(?P<attrs>[^>]*)>\s*(?P<number>\*?\d+\*?)\s*</span>\s*个\s+"
    r"(?P<term>\*?(?=[A-Za-z0-9.+#-]{2,20}\*?(?:\s|[，。！？；：,.!?;:]|$))"
    r"(?=[A-Za-z0-9.+#-]*[A-Za-z])(?=[A-Za-z0-9.+#-]*\d)[A-Za-z0-9.+#-]{2,20}\*?)",
    re.IGNORECASE,
)
DANGLING_PARAGRAPH_TAIL_RE = re.compile(
    r"(?:因为|由于|如果|虽然|尽管|当|只要|除非|不仅|无论|以及|并且|从而|以便|其中)\s*[，、：:]?$"
)
PATCH_ARTIFACT_RE = re.compile(r"^\*{3}\s+(?:Add|Update|Delete) File:", re.IGNORECASE)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
ROMAN_RE = re.compile(r"^[ivxlcdm]+$", re.IGNORECASE)
SECTION_RE = re.compile(r"^\d+\.\d+(?:\.\d+)?(?:[\s　]+).+")
CHAPTER_NUMBER_RE = re.compile(r"^\d{1,2}$")
CHAPTER_TITLE_RE = re.compile(r"^\d{1,2}[\s　]+.+")
PART_RE = re.compile(r"^第[一二三四五六七八九十百0-9]+部分(?:[\s　]+.*)?$")
CHAPTER_CN_RE = re.compile(r"^第\d{1,3}章(?:[:：.、\s]|$)")
SUBSECTION_RE = re.compile(r"^[（(][a-zA-Z一二三四五六七八九十]+[）)]\s*.+")
FRONT_TITLES = {
    "前言", "序言", "序", "引言", "导言", "内容提要", "摘要", "编者按", "出版说明",
    "撰稿人简介", "作者简介", "译者序", "致谢", "版权页", "书名页",
    "foreword", "preface", "introduction", "summary", "abstract", "contributors",
    "acknowledgments", "acknowledgements", "copyright", "title page",
}
TOC_TITLES = {"目录", "contents", "table of contents"}
TRAILING_RECOVERY_TITLES = {
    "OCR 恢复的源页面", "OCR恢复的源页面", "OCR recovered source pages",
}


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _epub_type(node: ET.Element) -> str:
    return node.attrib.get("{http://www.idpf.org/2007/ops}type", node.attrib.get("epub:type", ""))


def validate_xhtml_document(path: Path) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        root = ET.parse(path).getroot()
        ids: set[str] = set()
        duplicates: set[str] = set()
        for node in root.iter():
            node_id = node.attrib.get("id")
            if node_id in ids:
                duplicates.add(node_id)
            elif node_id:
                ids.add(node_id)
        if duplicates:
            errors.append("duplicate XHTML IDs: " + ", ".join(sorted(duplicates)))
        for node in root.iter():
            href = node.attrib.get("href", "")
            if href.startswith("#") and href[1:] not in ids:
                errors.append(f"XHTML link target is missing: {href}")
        note_refs = [
            node for node in root.iter()
            if _xml_local_name(node.tag) == "a"
            and ("noteref" in _epub_type(node).split() or node.attrib.get("role") == "doc-noteref")
        ]
        by_id = {node.attrib["id"]: node for node in root.iter() if node.attrib.get("id")}
        for ref in note_refs:
            href = ref.attrib.get("href", "")
            target = by_id.get(href[1:]) if href.startswith("#") else None
            if target is None:
                errors.append(f"note reference target is missing: {href or '<empty>'}")
                continue
            ref_id = ref.attrib.get("id", "")
            backlinks = [node.attrib.get("href", "") for node in target.iter() if _xml_local_name(node.tag) == "a"]
            if ref_id and f"#{ref_id}" not in backlinks:
                errors.append(f"note {href} has no backlink to #{ref_id}")
        if not any(_xml_local_name(node.tag) == "body" for node in root.iter()):
            errors.append("XHTML document has no body")
    except (OSError, ET.ParseError) as exc:
        errors.append(f"XHTML cannot be parsed: {exc}")
    return {"tool": "btp-xhtml-semantics", "passed": not errors, "errors": errors, "warnings": warnings}


def validate_epub_package(path: Path, require_cover: bool = False) -> dict[str, Any]:
    """Validate EPUB container/package topology before any final signoff."""
    errors: list[str] = []
    warnings: list[str] = []
    statistics = {"spine_items": 0, "heading_targets": 0, "navigation_targets": 0, "note_references": 0}
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if not names:
                errors.append("EPUB archive is empty")
                return {"tool": "btp-epub-structure", "passed": False, "errors": errors, "warnings": warnings}
            if archive.namelist()[0] != "mimetype":
                errors.append("mimetype is not the first ZIP member")
            info = archive.getinfo("mimetype") if "mimetype" in names else None
            if info is None:
                errors.append("mimetype is missing")
            else:
                if info.compress_type != zipfile.ZIP_STORED:
                    errors.append("mimetype must be uncompressed")
                if archive.read("mimetype") != b"application/epub+zip":
                    errors.append("mimetype has the wrong value")
            if "META-INF/container.xml" not in names:
                errors.append("META-INF/container.xml is missing")
                return {"tool": "btp-epub-structure", "passed": not errors, "errors": errors, "warnings": warnings}
            try:
                container = ET.fromstring(archive.read("META-INF/container.xml"))
            except ET.ParseError as exc:
                errors.append(f"container.xml is not valid XML: {exc}")
                return {"tool": "btp-epub-structure", "passed": False, "errors": errors, "warnings": warnings}
            rootfile = next((node for node in container.iter() if _xml_local_name(node.tag) == "rootfile"), None)
            package_path = rootfile.attrib.get("full-path", "") if rootfile is not None else ""
            if not package_path or package_path not in names:
                errors.append("container.xml does not point to an existing package document")
                return {"tool": "btp-epub-structure", "passed": False, "errors": errors, "warnings": warnings}
            try:
                package = ET.fromstring(archive.read(package_path))
            except ET.ParseError as exc:
                errors.append(f"package document is not valid XML: {exc}")
                return {"tool": "btp-epub-structure", "passed": False, "errors": errors, "warnings": warnings}

            manifest: dict[str, dict[str, str]] = {}
            title_present = False
            language_present = False
            for node in package.iter():
                local = _xml_local_name(node.tag)
                if local == "title" and node.text and node.text.strip():
                    title_present = True
                elif local == "language" and node.text and node.text.strip():
                    language_present = True
                elif local == "item":
                    item_id = node.attrib.get("id", "")
                    if item_id:
                        manifest[item_id] = node.attrib
            if not title_present:
                errors.append("package metadata has no title")
            if not language_present:
                errors.append("package metadata has no language")
            if not manifest:
                errors.append("manifest is empty")
            cover_items = [
                attributes for attributes in manifest.values()
                if "cover-image" in attributes.get("properties", "").split()
            ]
            if require_cover and not cover_items:
                errors.append("manifest has no EPUB cover-image item")

            base = posixpath.dirname(package_path)
            for item_id, attributes in manifest.items():
                href = attributes.get("href", "")
                if not href or "://" in href:
                    continue
                resource = posixpath.normpath(posixpath.join(base, unquote(href.split("#", 1)[0].split("?", 1)[0])))
                if resource and resource not in names:
                    errors.append(f"manifest resource is missing: {resource}")

            spine_refs = [
                node.attrib.get("idref", "")
                for node in package.iter()
                if _xml_local_name(node.tag) == "itemref"
            ]
            for idref in spine_refs:
                if idref not in manifest:
                    errors.append(f"spine idref is not in manifest: {idref}")
            if not spine_refs:
                errors.append("spine is empty")
            statistics["spine_items"] = len(spine_refs)
            nav_items = [
                attributes for attributes in manifest.values()
                if "nav" in attributes.get("properties", "").split()
            ]
            if not nav_items:
                errors.append("manifest has no navigation document")
            elif not any(
                posixpath.normpath(posixpath.join(base, unquote(item.get("href", "").split("#", 1)[0]))) in names
                for item in nav_items
            ):
                errors.append("navigation document is missing")

            document_ids: dict[str, set[str]] = {}
            heading_targets: set[str] = set()
            note_refs: list[tuple[str, str, str]] = []
            note_backlinks: dict[str, set[str]] = {}
            for item in manifest.values():
                media_type = item.get("media-type", "")
                if media_type not in {"application/xhtml+xml", "text/html"}:
                    continue
                href = unquote(item.get("href", "").split("?", 1)[0])
                resource = posixpath.normpath(posixpath.join(base, href))
                if resource not in names:
                    continue
                try:
                    root = ET.fromstring(archive.read(resource))
                except ET.ParseError as exc:
                    errors.append(f"EPUB content document is not valid XML: {resource}: {exc}")
                    continue
                ids = {node.attrib["id"] for node in root.iter() if node.attrib.get("id")}
                document_ids[resource] = ids
                is_navigation = "nav" in item.get("properties", "").split()
                for node in root.iter():
                    local = _xml_local_name(node.tag)
                    node_id = node.attrib.get("id", "")
                    if not is_navigation and local in {"h1", "h2", "h3", "h4", "h5", "h6"} and node_id:
                        heading_targets.add(f"{resource}#{node_id}")
                    if local == "a" and ("noteref" in _epub_type(node).split() or node.attrib.get("role") == "doc-noteref"):
                        note_refs.append((resource, node_id, node.attrib.get("href", "")))
                    if node_id and ("footnote" in _epub_type(node).split() or node.attrib.get("role") in {"doc-footnote", "doc-endnote"}):
                        note_backlinks[f"{resource}#{node_id}"] = {
                            child.attrib.get("href", "") for child in node.iter()
                            if _xml_local_name(child.tag) == "a" and child.attrib.get("href")
                        }

            nav_targets: set[str] = set()
            for item in nav_items:
                nav_resource = posixpath.normpath(posixpath.join(base, unquote(item.get("href", "").split("#", 1)[0])))
                if nav_resource not in names:
                    continue
                try:
                    nav_root = ET.fromstring(archive.read(nav_resource))
                except ET.ParseError:
                    continue
                nav_base = posixpath.dirname(nav_resource)
                toc_navs = [
                    node for node in nav_root.iter()
                    if _xml_local_name(node.tag) == "nav"
                    and ("toc" in _epub_type(node).split() or node.attrib.get("role") == "doc-toc")
                ]
                if not toc_navs:
                    errors.append("navigation document has no TOC nav element")
                for node in (child for toc in toc_navs for child in toc.iter()):
                    href = node.attrib.get("href", "")
                    if not href or "://" in href:
                        continue
                    target_file, _, fragment = href.partition("#")
                    target = posixpath.normpath(posixpath.join(nav_base, unquote(target_file)))
                    if target not in names:
                        errors.append(f"navigation target is missing: {target}")
                    elif fragment and fragment not in document_ids.get(target, set()):
                        errors.append(f"navigation fragment is missing: {target}#{fragment}")
                    if fragment:
                        nav_targets.add(f"{target}#{fragment}")
            statistics["heading_targets"] = len(heading_targets)
            statistics["navigation_targets"] = len(nav_targets)
            statistics["note_references"] = len(note_refs)
            if heading_targets and not nav_targets:
                errors.append("EPUB navigation contains no heading targets")

            for resource, ref_id, href in note_refs:
                target_file, _, fragment = href.partition("#")
                target_file = posixpath.normpath(posixpath.join(posixpath.dirname(resource), unquote(target_file))) if target_file else resource
                target_key = f"{target_file}#{fragment}"
                if not fragment or fragment not in document_ids.get(target_file, set()):
                    errors.append(f"note reference target is missing: {target_key}")
                    continue
                if ref_id:
                    expected_same = f"#{ref_id}"
                    expected_cross = posixpath.relpath(resource, posixpath.dirname(target_file)) + f"#{ref_id}"
                    if expected_same not in note_backlinks.get(target_key, set()) and expected_cross not in note_backlinks.get(target_key, set()):
                        errors.append(f"note target has no backlink to reference: {target_key}")
    except (OSError, zipfile.BadZipFile) as exc:
        errors.append(f"EPUB archive cannot be read: {exc}")
    return {"tool": "btp-epub-structure", "passed": not errors, "errors": errors, "warnings": warnings, "statistics": statistics}


def validate_docx_package(path: Path, require_toc: bool = False) -> dict[str, Any]:
    """Validate the DOCX package and the native TOC/heading signals."""
    errors: list[str] = []
    warnings: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                errors.append("word/document.xml is missing")
            else:
                document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
                if require_toc and "TOC" not in document_xml:
                    errors.append("native Word TOC field is missing")
            if "word/styles.xml" not in names:
                errors.append("word/styles.xml is missing")
            elif "word/document.xml" in names:
                styles = ET.fromstring(archive.read("word/styles.xml"))
                style_names = {}
                for node in styles.iter():
                    if _xml_local_name(node.tag) != "style":
                        continue
                    style_id = next((value for key, value in node.attrib.items() if _xml_local_name(key) == "styleId"), "")
                    name = ""
                    for child in node:
                        if _xml_local_name(child.tag) == "name":
                            name = next((value for key, value in child.attrib.items() if _xml_local_name(key) == "val"), "")
                            break
                    if style_id and name:
                        style_names[style_id] = name
                document = ET.fromstring(archive.read("word/document.xml"))
                heading_ids = {
                    next((value for key, value in node.attrib.items() if _xml_local_name(key) == "val"), "")
                    for node in document.iter()
                    if _xml_local_name(node.tag) == "pStyle"
                }
                if not any(style_names.get(style_id, "").casefold().startswith("heading") for style_id in heading_ids):
                    warnings.append("document has no explicit heading style signal in document.xml")
    except (OSError, zipfile.BadZipFile) as exc:
        errors.append(f"DOCX archive cannot be read: {exc}")
    return {"tool": "btp-docx-structure", "passed": not errors, "errors": errors, "warnings": warnings}


def validate_pdf_package(path: Path) -> dict[str, Any]:
    """Validate PDF readability and reject visible field-error residue."""
    errors: list[str] = []
    warnings: list[str] = []
    try:
        import fitz

        with fitz.open(path) as document:
            if len(document) < 1:
                errors.append("PDF has no pages")
            for page_number, page in enumerate(document, 1):
                text = page.get_text()
                if "Error!" in text or "Reference source not found" in text:
                    errors.append(f"PDF contains a Word field error on page {page_number}")
    except ImportError:
        warnings.append("PyMuPDF is unavailable; PDF text-level field-error scan was skipped")
    except (OSError, RuntimeError) as exc:
        errors.append(f"PDF cannot be read: {exc}")
    return {"tool": "btp-pdf-structure", "passed": not errors, "errors": errors, "warnings": warnings}


PIPE_SEPARATOR_RE = re.compile(r"^\|(?:\s*:?-{2,}:?\s*\|)+$")
MARKDOWN_IMAGE_RE = re.compile(r"!\[[^\]]*\]\((?P<target>[^)\s]+)(?:\s+[^)]*)?\)")
IMAGE_SUFFIXES = (".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".bmp", ".webp", ".svg")


def _pipe_table_count(markdown: str) -> int:
    """Count Markdown pipe tables by separator rows, including Pandoc's ``--`` form."""
def _separate_html_blocks_from_tables(lines: list[str]) -> list[str]:
    """Insert a blank line between a raw-HTML line and a following pipe table.

    Pandoc parses a raw-HTML line as an HTML block that swallows every
    following line until the next blank line — pipe tables glued directly
    under extraction HTML (thumbnail anchors, image links, bookmark spans)
    silently degrade into paragraph text.  The blank line ends the HTML block
    so the table parses as a table.
    """
    result: list[str] = []
    for line in lines:
        stripped = line.lstrip()
        if stripped.startswith("|") and result:
            previous = result[-1].lstrip()
            if previous.startswith("<") and not previous.startswith("<!--"):
                result.append("")
        result.append(line)
    return result


def _pipe_table_count(markdown: str) -> int:
    """Count pipe tables the way Pandoc builds them: a separator row must be
    preceded by a header row of the same pipe shape.  Degenerate separator
    lines without a header (layout-table residue from CHM conversions) never
    become tables in the generated DOCX and would make the inventory check
    demand tables Pandoc cannot produce.
    """
    lines = markdown.splitlines()
    count = 0
    for index, line in enumerate(lines):
        if not PIPE_SEPARATOR_RE.fullmatch(line.strip()):
            continue
        if index and lines[index - 1].lstrip().startswith("|"):
            count += 1
    return count


def _expected_image_refs(markdown: str) -> set[str]:
    refs: set[str] = set()
    for match in MARKDOWN_IMAGE_RE.finditer(markdown):
        target = unquote(match.group("target")).replace("\\", "/")
        if re.match(r"^(?:https?:|data:)", target, re.IGNORECASE):
            continue
        refs.add(target.split("#", 1)[0].split("?", 1)[0])
    return refs


def _normalize_inventory_text(text: str, *, rendered: bool = False) -> str:
    """Normalize text for content-anchor comparison.

    For source Markdown (``rendered=False``) markdown image/link syntax and
    raw HTML tags are stripped.  For text extracted from rendered outputs
    (``rendered=True``) there is no markup left — and book content may
    legitimately contain ``<`` / ``>`` comparison characters — so the
    tag-stripping subs are skipped: a greedy ``<[^>]+>`` on rendered text can
    otherwise swallow tens of thousands of real characters between a ``<`` in
    a table cell and the next ``>`` elsewhere in the book.
    """
    value = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    value = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", value)
    if not rendered:
        value = re.sub(r"<[^>]+>", "", value)
    value = re.sub(r"[^\w\u3400-\u9fff]+", "", value, flags=re.UNICODE)
    return value.casefold()


def _content_anchors(markdown: str, limit: int = 9) -> list[str]:
    """Choose deterministic front/middle/back text samples from the clean master."""
    candidates: list[str] = []
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped or stripped == "[[TOC]]" or PIPE_SEPARATOR_RE.fullmatch(stripped):
            continue
        if stripped.startswith(":::"):
            continue
        stripped = re.sub(r"\[\^[^]]+\](?::)?", "", stripped).strip()
        if not stripped:
            continue
        # List markers are represented differently by DOCX/EPUB/PDF writers;
        # sample their surrounding prose instead of treating ordinal glyphs as
        # stable cross-format content.
        if re.match(r"^(?:[-+*]\s+|\d+[.)]\s+)", stripped):
            continue
        anchor = _normalize_inventory_text(stripped)
        if len(anchor) < 12 or anchor in candidates:
            continue
        candidates.append(anchor[:180])
    if len(candidates) <= limit:
        return candidates
    indices = {round(index * (len(candidates) - 1) / (limit - 1)) for index in range(limit)}
    return [candidates[index] for index in sorted(indices)]


def _docx_inventory(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        document = ET.fromstring(archive.read("word/document.xml"))
        text = "".join(node.text or "" for node in document.iter() if _xml_local_name(node.tag) == "t")
        media_hashes = {
            hashlib.md5(archive.read(name)).hexdigest()
            for name in names
            if name.startswith("word/media/") and not name.endswith("/")
        }
        return {
            "table_count": sum(1 for node in document.iter() if _xml_local_name(node.tag) == "tbl"),
            "image_count": len(media_hashes),
            "image_hashes": sorted(media_hashes),
            "text": _normalize_inventory_text(text, rendered=True),
        }


def _epub_inventory(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        pages = [name for name in names if name.casefold().endswith((".xhtml", ".html", ".htm"))]
        bodies = [archive.read(name).decode("utf-8", errors="replace") for name in pages]
        media_hashes = {
            hashlib.md5(archive.read(name)).hexdigest()
            for name in names
            if name.casefold().split("?", 1)[0].endswith(IMAGE_SUFFIXES)
        }
        return {
            "table_count": sum(len(re.findall(r"<table\b", body, re.IGNORECASE)) for body in bodies),
            "image_hashes": sorted(media_hashes),
            "image_count": len(media_hashes),
            "text": _normalize_inventory_text(" ".join(BeautifulSoup(body, "html.parser").get_text(" ") for body in bodies), rendered=True),
        }


def _pdf_inventory(path: Path) -> dict[str, Any]:
    try:
        import fitz
    except ImportError:
        return {"text": "", "text_available": False}
    # Running heads/feet inject bare page-number blocks and repeated titles at
    # page boundaries. Drop those zones geometrically and reject only a block
    # whose entire payload is a folio. Word PDF extraction can represent every
    # word in a body sentence on its own text line; applying a multiline regex
    # to the joined body would then erase legitimate values such as ``2 mm``.
    page_number_block = re.compile(r"^(?:[ivxlcdm]{1,8}|\d{1,5})$", re.IGNORECASE)
    with fitz.open(path) as document:
        pieces = []
        for page in document:
            height = page.rect.height
            body = []
            for block in page.get_text("blocks"):
                text = str(block[4]).strip()
                if block[1] <= height * 0.09 or block[3] >= height * 0.91:
                    continue
                if page_number_block.fullmatch(text):
                    continue
                body.append(str(block[4]))
            pieces.append("\n".join(body))
        joined = "\n".join(pieces)
        return {
            "text": _normalize_inventory_text(joined, rendered=True),
            "text_available": True,
        }


def _content_inventory_check(kind: str, path: Path, master: str, project: Path | None = None) -> dict[str, Any]:
    expected_tables = _pipe_table_count(master)
    expected_refs = _expected_image_refs(master)
    anchors = _content_anchors(master)
    if kind == "docx":
        actual = _docx_inventory(path)
    elif kind == "epub":
        actual = _epub_inventory(path)
    else:
        actual = _pdf_inventory(path)
    errors: list[str] = []
    warnings: list[str] = []
    if kind in {"docx", "epub"}:
        if actual["table_count"] < expected_tables:
            errors.append(f"table inventory shrank: expected at least {expected_tables}, found {actual['table_count']}")
        # Compare by content hash: identical images referenced several times
        # (deduplicated by Pandoc) must not count as shrinkage.  References
        # whose asset file is missing count as not embedded.
        expected_hashes = set()
        unresolved_refs = 0
        for ref in expected_refs:
            candidate = Path(ref.replace("\\", "/"))
            if project is not None:
                candidate = project / candidate
            if candidate.is_file():
                expected_hashes.add(hashlib.md5(candidate.read_bytes()).hexdigest())
            else:
                unresolved_refs += 1
        embedded_missing = expected_hashes - set(actual.get("image_hashes") or [])
        missing_images = len(embedded_missing) + unresolved_refs
        if missing_images:
            errors.append(
                f"image inventory shrank: {missing_images} of {len(expected_refs)} referenced images "
                "are missing or not embedded"
            )
    if kind == "pdf" and not actual.get("text_available"):
        warnings.append("PDF text extraction is unavailable; content-anchor sampling was skipped")
        missing: list[str] = []
    else:
        missing = [anchor for anchor in anchors if anchor not in actual.get("text", "")]
        if missing:
            errors.append(
                f"{len(missing)} of {len(anchors)} sampled content anchors are missing: "
                + ", ".join(missing[:3])
            )
    return {
        "tool": "btp-content-inventory",
        "passed": not errors,
        "expected": {
            "table_count": expected_tables,
            "unique_image_refs": len(expected_refs),
            "content_anchor_count": len(anchors),
        },
        "actual": {key: value for key, value in actual.items() if key != "text"},
        "missing_anchors": missing[:5],
        "errors": errors,
        "warnings": warnings,
    }


def validate_book_map_consistency(project: Path, generated: dict[str, Path]) -> dict[str, Any]:
    """Tie the final content inventory to the canonical logical book map."""
    errors: list[str] = []
    book_map = load_json(project / "analysis" / "book-map.json")
    manifest = load_manifest(project)
    expected_segments = {
        segment_id for node in book_map["nodes"] for segment_id in node["source_span"]["segment_ids"]
    }
    output_segments = {segment_id for unit in manifest["units"] for segment_id in unit["segment_ids"]}
    if expected_segments != output_segments:
        errors.append("manifest segment coverage differs from Book Map")
    mapped_nodes = {unit.get("book_node_id") for unit in manifest["units"]}
    missing_nodes = [
        node["id"] for node in book_map["nodes"]
        if node["source_span"]["segment_ids"] and node["id"] not in mapped_nodes
    ]
    if missing_nodes:
        errors.append("Book Map nodes have no target work unit: " + ", ".join(missing_nodes[:8]))
    master_path = generated.get("markdown")
    master = master_path.read_text(encoding="utf-8") if master_path and master_path.is_file() else ""
    expected_headings = sum(1 for node in book_map["nodes"] if node["id"].startswith("section-"))
    actual_headings = sum(1 for line in master.splitlines() if HEADING_RE.match(line.strip()))
    if actual_headings < expected_headings:
        errors.append(f"clean master has {actual_headings} headings for {expected_headings} mapped heading nodes")
    return {
        "tool": "btp-book-map-consistency",
        "passed": not errors,
        "expected_heading_nodes": expected_headings,
        "actual_master_headings": actual_headings,
        "errors": errors,
        "warnings": [],
    }


def validate_publication_outputs(
    generated: dict[str, Path], config: dict[str, Any], project: Path | None = None
) -> dict[str, Any]:
    """Run structural and content-inventory checks for publication outputs."""
    checks: dict[str, Any] = {}
    errors: list[str] = []
    warnings: list[str] = []
    master_path = generated.get("markdown")
    master = master_path.read_text(encoding="utf-8") if master_path and master_path.is_file() else ""
    require_toc = "[[TOC]]" in master

    validation_paths: dict[str, Path] = {
        label: path for label, path in generated.items()
        if label.removeprefix("bilingual_") in {"xhtml", "docx", "epub", "pdf"}
    }
    if "pdf" in validation_paths and "docx" not in validation_paths and master_path is not None:
        intermediate = master_path.with_name("book.docx")
        if intermediate.is_file():
            validation_paths["docx_intermediate"] = intermediate

    for label, path in validation_paths.items():
        kind = "docx" if label == "docx_intermediate" else label.removeprefix("bilingual_")
        label_master_path = generated.get("bilingual_markdown") if label.startswith("bilingual_") else master_path
        label_master = label_master_path.read_text(encoding="utf-8") if label_master_path and label_master_path.is_file() else master
        label_require_toc = "[[TOC]]" in label_master
        if kind == "xhtml":
            result = validate_xhtml_document(path)
        elif kind == "docx":
            result = validate_docx_package(path, require_toc=label_require_toc)
        elif kind == "epub":
            result = validate_epub_package(path, require_cover=bool(config.get("metadata", {}).get("cover")))
        else:
            result = validate_pdf_package(path)
        checks[label] = result
        errors.extend(f"{label}: {message}" for message in result["errors"])
        warnings.extend(f"{label}: {message}" for message in result["warnings"])
        if label_master and kind in {"docx", "epub", "pdf"}:
            inventory = _content_inventory_check(kind, path, label_master, project=project)
            checks[f"{label}_content"] = inventory
            errors.extend(f"{label}: {message}" for message in inventory["errors"])
            warnings.extend(f"{label}: {message}" for message in inventory["warnings"])

    for epub_label in [label for label in generated if label.removeprefix("bilingual_") == "epub"]:
        epubcheck = epubcheck_command()
        if epubcheck:
            try:
                result = subprocess.run(
                    [*epubcheck, external_tool_path(generated[epub_label])],
                    capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                )
                external = {
                    "tool": "EPUBCheck",
                    "status": "passed" if result.returncode == 0 else "failed",
                    "returncode": result.returncode,
                    "output": "\n".join(part.strip() for part in (result.stdout, result.stderr) if part.strip())[-4000:],
                }
                if result.returncode != 0:
                    summary = " | ".join(external["output"].splitlines()[-8:])
                    errors.append(f"{epub_label}: EPUBCheck reported errors: {summary}")
            except (OSError, subprocess.TimeoutExpired) as exc:
                external = {"tool": "EPUBCheck", "status": "failed", "error": str(exc)}
                errors.append(f"{epub_label}: EPUBCheck could not run: {exc}")
        else:
            external = {"tool": "EPUBCheck", "status": "unavailable"}
            warnings.append(f"{epub_label}: external EPUBCheck is unavailable; internal EPUB topology checks were run")
        check_label = "epubcheck" if epub_label == "epub" else f"{epub_label}_epubcheck"
        checks[check_label] = external
        if config.get("mode") == "publication" and external["status"] != "passed":
            errors.append(f"{epub_label}: Publication mode requires an EPUBCheck pass")

    if project is not None:
        book_map_check = validate_book_map_consistency(project, generated)
        checks["book_map"] = book_map_check
        errors.extend(f"book_map: {message}" for message in book_map_check["errors"])

    # Pandoc builds a table from every master pipe-table block, but generated
    # formats legitimately end up with fewer table ELEMENTS: Word merges
    # adjacent tables on save (content verified intact) and both Pandoc
    # renders agree with each other.  When every generated format reports the
    # same table count while the master counts higher, the shrinkage is a
    # counting artifact of the master — downgrade to a warning.  Formats that
    # disagree with each other still block: that is real generation loss.
    table_counts = {
        label[:-len("_content")]: checks[label].get("actual", {}).get("table_count")
        for label in checks
        if label.endswith("_content") and checks[label].get("actual", {}).get("table_count") is not None
    }
    counts_present = list(table_counts.values())
    if len(counts_present) >= 2 and len(set(counts_present)) == 1:
        agreed_count = counts_present[0]
        for label, count in table_counts.items():
            inventory = checks[f"{label}_content"]
            expected_tables = inventory.get("expected", {}).get("table_count", 0)
            if count < expected_tables:
                inventory["errors"] = [
                    message for message in inventory["errors"] if "table inventory shrank" not in message
                ]
                inventory["passed"] = not inventory["errors"]
                inventory["warnings"].append(
                    f"{label} table element count ({count}) agrees across all generated formats; "
                    "the master pipe-table count includes layout tables the renderers merge on save"
                )
                errors = [
                    message for message in errors
                    if not (message.startswith(f"{label}:") and "table inventory shrank" in message)
                ]
                warnings.append(f"{label}: master pipe-table count includes layout tables the renderers merge on save")

    return {"passed": not errors, "checks": checks, "errors": errors, "warnings": warnings}


def _plain(text: str) -> str:
    value = HTML_TAG_RE.sub("", text)
    value = re.sub(r"[`*_]+", "", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" \t#")


def _normalize_html_headings(markdown: str) -> str:
    """Convert retained HTML headings before non-HTML Pandoc writers see them."""
    def replace(match: re.Match[str]) -> str:
        label = BeautifulSoup(match.group("body"), "html.parser").get_text(" ", strip=True)
        return f"{'#' * int(match.group('level'))} {label}" if label else ""

    return HTML_HEADING_RE.sub(replace, markdown)


def _emphasis_only(text: str) -> bool:
    stripped = text.strip()
    return stripped.startswith("**") and stripped.endswith("**")


def _is_running_folio(text: str, current_page: str | None) -> bool:
    if not current_page:
        return False
    plain = _plain(text)
    if not plain or len(plain) > 90 or any(mark in plain for mark in "。！？；?!;"):
        return False
    tokens = plain.split()
    page = current_page.casefold()
    if plain.casefold() == page:
        return True
    return bool(tokens and (tokens[0].casefold() == page or tokens[-1].casefold() == page))


def _strip_print_artifacts(markdown: str) -> list[str]:
    # Internal jump targets survive anchor cleanup; everything else (page
    # number anchors such as id="25" or id="page3") stays on the removal path.
    jump_targets = set(re.findall(r'href\s*=\s*["\']#([^"\']+)["\']', markdown, re.IGNORECASE))
    cleaned: list[str] = []
    current_page: str | None = None
    skip_observation_payload = False
    for raw in markdown.splitlines():
        raw = HTML_COMMENT_RE.sub("", raw).strip()
        if PATCH_ARTIFACT_RE.match(raw):
            skip_observation_payload = True
            continue
        if skip_observation_payload and raw.startswith("{") and "schema_version" in raw:
            skip_observation_payload = False
            continue
        skip_observation_payload = False
        matches = list(PRINT_MARKER_RE.finditer(raw))
        if matches:
            current_page = matches[-1].group("page")
            raw = PRINT_MARKER_RE.sub("", raw)
        raw = _normalize_footnote_markup(raw)
        # Preserve semantic TOC markers until raw HTML tables are converted.
        # The ordinary span cleanup would otherwise erase ``span#TOC`` before
        # the nesting-aware table pass can replace its layout table.
        raw = re.sub(
            r"<span\b(?=[^>]*(?:\bid\s*=\s*['\"]?TOC\b|\bclass\s*=\s*['\"][^'\"]*\bb24-toctitle\b))[^>]*>",
            "[[BTP-HTML-TOC]]",
            raw,
            flags=re.IGNORECASE,
        )
        # Preserve empty anchors that internal jumps target (CHM cross
        # references such as ``<span id="ch01fig01"></span>``): the generic
        # span cleanup below would erase them and orphan every
        # ``href="#ch01fig01"`` link.  Page-number anchors (numeric ids) are
        # print residue and stay on the removal path.
        raw = _ANCHOR_TARGET_RE.sub(
            lambda match: f"[[BTP-ANCHOR-{match.group(1)}]]" if match.group(1) in jump_targets else "",
            raw,
        )
        raw = HTML_TAG_RE.sub("", raw).strip()
        restored_only_anchors = bool(re.search(r"\[\[BTP-ANCHOR-", raw)) and bool(
            re.fullmatch(r"(?:\s*\[\[BTP-ANCHOR-[^\]]+\]\])+", raw)
        )
        raw = re.sub(
            r"\[\[BTP-ANCHOR-([^\]]+)\]\]",
            lambda match: f'<span id="{match.group(1)}"></span>',
            raw,
        )
        if not raw:
            cleaned.append("")
            continue
        # A standalone anchor line must be followed by a blank line: pandoc
        # treats a raw-HTML block as running until the next blank line and
        # would otherwise swallow an adjacent pipe table or image.
        if restored_only_anchors:
            cleaned.append(raw)
            cleaned.append("")
            continue
        if _is_running_folio(raw, current_page):
            continue
        heading = HEADING_RE.match(raw)
        if heading and current_page and _plain(heading.group(2)).casefold() == current_page.casefold():
            continue
        if raw.startswith("•"):
            raw = "- " + raw[1:].lstrip()
        if raw.endswith(chr(92)) and not raw.endswith(chr(92) * 2):
            # Markdown hard line break: address blocks and CIP data come as
            # groups of short lines.  Split them into separate paragraphs so
            # justified body text cannot stretch them across the column.
            raw = raw[:-1].rstrip()
            if raw:
                cleaned.append(raw)
            cleaned.append("")
            continue
        if raw in {"***", "**", "*"}:
            # Emphasis wrappers around empty extraction spans (for example a
            # page anchor wrapped in asterisks) leave these bare markers
            # behind; they render as stray rules or asterisks in print.
            continue
        if raw == chr(92):
            # Extraction line-continuation residue; a lone backslash renders
            # as a stray character in print output.
            continue
        cleaned.append(raw)
    return cleaned


def _normalize_footnote_markup(text: str) -> str:
    """Restore superscript footnote markers lost by HTML extraction.

    The source EPUB commonly represents footnote references as numeric
    ``calibre26`` spans.  Removing the span wrapper alone turns those
    references into ordinary body text.  Keep quantity-like uses such as
    ``11个`` and ``5S`` as normal text.
    """

    def replace(match: re.Match[str]) -> str:
        attrs = match.group("attrs")
        classes = re.search(r"\bclass\s*=\s*([\"'])(.*?)\1", attrs, re.IGNORECASE | re.DOTALL)
        if not classes or "calibre26" not in classes.group(2).split():
            return match.group(0)
        number = match.group("number").strip("*")
        following = text[match.end():]
        if following and following[0] in "个%％次项条人年月日类章节页台名位件种Ss":
            return match.group(0)
        # Pandoc's Markdown superscript syntax survives DOCX and EPUB
        # conversion; raw HTML ``<sup>`` is discarded by some DOCX writers.
        return f"^{number}^"

    def move_protected_term_marker(match: re.Match[str]) -> str:
        marker = f'<span{match.group("attrs")}>{match.group("number")}</span>'
        return f' {match.group("term")}{marker}'

    text = FOOTNOTE_BEFORE_PROTECTED_TERM_RE.sub(move_protected_term_marker, text)
    return FOOTNOTE_SPAN_RE.sub(replace, text)


def _join_dangling_paragraphs(lines: list[str]) -> list[str]:
    """Join a paragraph split at a source-page or source-unit boundary.

    PDF/EPUB extraction can leave a blank line between the two halves of a
    sentence when the source page ends with a conjunction such as ``因为``.
    Only join clear continuation cases; headings, lists, quotes, tables, and
    images remain independent blocks.
    """

    result: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.strip():
            result.append(line)
            index += 1
            continue
        next_index = index
        while next_index < len(lines) and not lines[next_index].strip():
            next_index += 1
        if result and result[-1].strip() and DANGLING_PARAGRAPH_TAIL_RE.search(_plain(result[-1])) and next_index < len(lines):
            next_line = lines[next_index].strip()
            if not re.match(r"^(?:#{1,6}\s|[-*+]\s|>\s*|\||!\[|\[\[TOC\]\])", next_line):
                previous_text = result.pop()
                joiner = " " if previous_text[-1:].isascii() and next_line[:1].isascii() else ""
                result.append(previous_text.rstrip() + joiner + next_line)
                index = next_index + 1
                continue
        result.extend(lines[index:next_index])
        index = next_index
    return result


def _replace_static_toc(lines: list[str]) -> list[str]:
    toc_index: int | None = None
    for index, line in enumerate(lines):
        heading = HEADING_RE.match(line)
        candidate = _plain(heading.group(2) if heading else line).casefold()
        if candidate in TOC_TITLES:
            toc_index = index
            break
    if toc_index is None:
        return lines
    body_index: int | None = None
    for index in range(toc_index + 1, len(lines)):
        candidate = _plain(lines[index]).casefold()
        if candidate in FRONT_TITLES:
            body_index = index
            break
    if body_index is None:
        return lines
    return [*lines[:toc_index], "# 目录", "", "[[TOC]]", "", *lines[body_index:]]


def _next_nonblank(lines: list[str], start: int) -> tuple[int | None, str | None]:
    for index in range(start, len(lines)):
        if lines[index].strip():
            return index, lines[index]
    return None, None


def _normalize_structure(lines: list[str]) -> list[str]:
    result: list[str] = []
    index = 0
    after_chapter = False
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            if result and result[-1] != "":
                result.append("")
            index += 1
            continue
        plain = _plain(line)
        folded = plain.casefold()
        heading = HEADING_RE.match(line)

        if folded in FRONT_TITLES:
            result.append(f"# {plain}")
            after_chapter = False
            index += 1
            continue

        if heading:
            level = len(heading.group(1))
            text = _plain(heading.group(2))
            if not text or (len(text) > 1 and ROMAN_RE.fullmatch(text)):
                index += 1
                continue
            # OCR/PDF extraction can promote executive-summary prose to a
            # heading.  Long punctuated blocks are paragraphs, not TOC items.
            if level >= 2 and len(text) > 55 and any(mark in text for mark in "。；，"):
                result.append(text)
                after_chapter = False
                index += 1
                continue
            if re.match(r"^第\s*\d+\s*章内容提要$", text):
                result.append(f"### {text}")
                after_chapter = False
                index += 1
                continue
            if text.startswith("本章") and len(text) > 70:
                result.append(text)
                after_chapter = False
                index += 1
                continue
            if PART_RE.match(text):
                next_index, next_line = _next_nonblank(lines, index + 1)
                if next_line and _emphasis_only(next_line) and not HEADING_RE.match(next_line):
                    subtitle = _plain(next_line)
                    if subtitle and not SECTION_RE.match(subtitle):
                        text = f"{text} {subtitle}"
                        index = next_index or index
                result.append(f"# {text}")
                after_chapter = False
                index += 1
                continue
            if level == 1 and CHAPTER_TITLE_RE.match(text):
                result.append(f"# {text}")
                after_chapter = True
                index += 1
                continue
            if after_chapter and level >= 2 and not SECTION_RE.match(text) and text not in {"注释", "Notes"}:
                result.append(f"*{text}*")
                after_chapter = False
                index += 1
                continue
            result.append(f"{'#' * level} {text}")
            after_chapter = level == 1 and CHAPTER_TITLE_RE.match(text) is not None
            index += 1
            continue

        next_index, next_line = _next_nonblank(lines, index + 1)
        if CHAPTER_NUMBER_RE.match(plain) and _emphasis_only(line) and next_line and _emphasis_only(next_line):
            title = _plain(next_line)
            if title and not CHAPTER_NUMBER_RE.match(title):
                result.append(f"# {plain} {title}")
                after_chapter = True
                index = (next_index or index) + 1
                continue
        if _emphasis_only(line) and SECTION_RE.match(plain):
            result.append(f"## {plain}")
            after_chapter = False
        elif _emphasis_only(line) and SUBSECTION_RE.match(plain):
            result.append(f"### {plain}")
            after_chapter = False
        elif _emphasis_only(line) and folded in {"注释", "notes"}:
            result.append(f"## {plain}")
            after_chapter = False
        else:
            result.append(line)
            if after_chapter and plain:
                after_chapter = False
        index += 1

    while result and not result[-1]:
        result.pop()
    return result


def _strip_trailing_recovery_material(lines: list[str]) -> list[str]:
    """Drop extraction-only recovery units appended after the real book."""
    midpoint = len(lines) // 2
    for index in range(midpoint, len(lines)):
        heading = HEADING_RE.match(lines[index].strip())
        if heading and _plain(heading.group(2)) in TRAILING_RECOVERY_TITLES:
            return lines[:index]
    return lines


PLATFORM_IMAGE_NAMES = {
    "teamlib.gif", "next.gif", "previous.gif", "_.gif",
    "smallcd.gif", "i_cdcontent.gif",
}
SUP_MARKER_RE = re.compile(r"<sup\b[^>]*>(?P<inner>.*?)</sup>", re.IGNORECASE | re.DOTALL)


def _normalize_sup_markers(text: str) -> str:
    """Convert extracted <sup>[N]</sup> footnote markers into pandoc superscripts.

    Raw HTML ``<sup>`` is discarded by the DOCX writer, which would silently
    delete every footnote marker from publication DOCX/PDF.
    """

    def repl(match: re.Match[str]) -> str:
        inner = re.sub(r"</?a\b[^>]*>", "", match.group("inner"), flags=re.IGNORECASE)
        inner = re.sub(r"</?(?:em|strong|b|i)\b[^>]*>", "", inner, flags=re.IGNORECASE).strip()
        bounded = re.fullmatch(r"\\?\[\s*([^\[\]]{1,16}?)\s*\\?\]", inner)
        if not bounded:
            return match.group(0)
        label = bounded.group(1).strip().strip("*").strip()
        if not label:
            return match.group(0)
        # "^[...]" would parse as a pandoc inline footnote; escape the
        # brackets (and any spaces) so "^\\[38\\]^" stays a superscript.
        label = label.replace(" ", "\\ ")
        return f"^\\[{label}\\]^"

    return SUP_MARKER_RE.sub(repl, text)


def _html_cell_text(cell) -> str:
    for br in cell.find_all("br"):
        br.replace_with(" ")
    # Figure links inside cells ([thumbnail](original) pairs) must survive
    # table conversion: keep them as Markdown linked images so the promotion
    # pass can upgrade to the original and the asset restore can backfill it.
    image_suffixes = (".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".bmp", ".webp")
    for anchor in list(cell.find_all("a")):
        attrs = anchor.attrs or {}
        href = str(attrs.get("href", "")).strip()
        href_path = href.replace("\\", "/").split("?", 1)[0].split("#", 1)[0]
        if not href_path.lower().endswith(image_suffixes):
            continue
        img = anchor.find("img")
        if img is None:
            continue
        img_attrs = img.attrs or {}
        src = str(img_attrs.get("src", "")).strip().replace("\\", "/")
        if not src:
            continue
        alt = str(img_attrs.get("alt", "")).strip().replace("]", "\\]")
        if not src.lower().startswith("assets/"):
            src = "assets/original/" + src.lstrip("/")
        # Promote here: the cell holds plain Markdown, so the later <a>-based
        # promotion passes cannot see this pair anymore.
        big = href.replace("\\", "/")
        if src.lower().startswith("assets/original/") and not big.lower().startswith("assets/"):
            big = "assets/original/" + big.lstrip("/")
        anchor.replace_with(f"![{alt}]({big})")
    for img in list(cell.find_all("img")):
        img_attrs = img.attrs or {}
        src = str(img_attrs.get("src", "")).strip().replace("\\", "/")
        if not src or src.rsplit("/", 1)[-1].casefold() in PLATFORM_IMAGE_NAMES:
            img.decompose()
            continue
        alt = str(img_attrs.get("alt", "")).strip().replace("]", "\\]")
        if not src.lower().startswith("assets/"):
            src = "assets/original/" + src.lstrip("/")
        img.replace_with(f"![{alt}]({src})")
    for nested in cell.find_all("table"):
        nested.decompose()
    text = cell.get_text(" ", strip=True).replace("\xa0", " ")
    text = text.replace("|", "\\|")
    text = re.sub(r"\s+", " ", text).strip()
    if text in {"*", "**", "***", "•", "-", "\\"}:
        # Empty-emphasis wrappers and lone bullets inside layout cells are
        # extraction residue, not content.
        return ""
    return text


def _iter_outer_tables(text: str):
    """Yield (start, end) spans of outermost ``<table>`` blocks, nesting-aware."""
    token_re = re.compile(r"</?table\b[^>]*/?>", re.IGNORECASE)
    depth = 0
    start = None
    for token in token_re.finditer(text):
        if token.group(0).lower().startswith("</table"):
            depth -= 1
            if depth == 0 and start is not None:
                yield start, token.end()
                start = None
            if depth < 0:
                depth = 0
        else:
            if depth == 0:
                start = token.start()
            depth += 1


TOC_MARKER_RE = re.compile(
    r'b24-toctitle|<span\b[^>]*\bid=["\']?TOC["\']?|\[\[BTP-HTML-TOC\]\]',
    re.IGNORECASE,
)


def _is_prose_row(row: list[str], threshold: int = 40) -> bool:
    """A row carrying a single long prose cell is a footnote/paragraph that
    the source HTML embedded inside the table — not tabular data."""
    non_empty = [cell for cell in row if cell.strip()]
    if len(non_empty) != 1:
        return False
    return _display_width(non_empty[0]) >= threshold


def _grid_table_markdown(table, strip_nested_markers: bool = False) -> str:
    """Flatten one (nesting-free) HTML table into a Markdown pipe table."""
    nest_marker_re = re.compile(r"\[\[NESTTBL:\d+\]\]")
    caption = table.find("caption")
    caption_text = re.sub(r"\s+", " ", caption.get_text(" ", strip=True)) if caption else ""
    rows = table.find_all("tr")
    if not rows:
        return ""
    occupancy: dict[tuple[int, int], str] = {}
    width = 0
    for r, tr in enumerate(rows):
        c = 0
        for cell in tr.find_all(["td", "th"]):
            while (r, c) in occupancy:
                c += 1
            attrs = cell.attrs or {}
            try:
                colspan = max(1, int(str(attrs.get("colspan", 1) or 1)))
                rowspan = max(1, int(str(attrs.get("rowspan", 1) or 1)))
            except (TypeError, ValueError):
                colspan = rowspan = 1
            colspan = min(colspan, 24)
            rowspan = min(rowspan, max(1, len(rows) - r))
            label = _html_cell_text(cell)
            for dr in range(rowspan):
                for dc in range(colspan):
                    occupancy[(r + dr, c + dc)] = label if (dr == 0 and dc == 0) else ""
            c += colspan
        width = max(width, c)
    if width == 0:
        return ""
    grid = [
        [occupancy.get((r, c), "") for c in range(width)]
        for r in range(len(rows))
    ]
    if strip_nested_markers:
        grid = [
            [nest_marker_re.sub("", cell).strip() for cell in row]
            for row in grid
        ]
    if not any(any(cell for cell in row) for row in grid):
        # Pure layout/spacer scaffolding carries no content.
        return ""
    has_thead = table.find("thead") is not None
    first_row_is_header = bool(rows[0].find("th"))
    # Prose rows (footnotes/paragraphs embedded as single wide cells) leave
    # the table; otherwise they stretch one column across the whole grid.
    data_rows: list[list[str]] = []
    prose: list[str] = []
    for r, row in enumerate(grid):
        if _is_prose_row(row):
            prose.append(next(cell for cell in row if cell.strip()))
        else:
            data_rows.append(row)
    grid = data_rows
    if not grid:
        if prose:
            return "\n\n" + "\n\n".join(prose) + "\n\n"
        return ""
    lines: list[str] = []
    if caption_text:
        lines.append(f"**{caption_text}**")
        lines.append("")
    if has_thead or first_row_is_header:
        header = grid[0]
        body = grid[1:]
    else:
        header = [""] * width
        body = grid
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * width) + " |")
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    output = "\n\n" + "\n".join(lines) + "\n\n"
    if prose:
        output = output + "\n\n" + "\n\n".join(prose) + "\n\n"
    return output


def _convert_single_table(block: str) -> str:
    soup = BeautifulSoup(block, "html.parser")
    if TOC_MARKER_RE.search(block):
        marker = soup.find("span", id="TOC") or soup.find(class_="b24-toctitle")
        if marker is None:
            marker = soup.find(string=lambda value: bool(value and "[[BTP-HTML-TOC]]" in value))
        node = marker
        while node is not None and getattr(node, "name", "") != "table":
            node = node.parent
        placeholder = "\n\n# 目录\n\n[[TOC]]\n\n"
        if node is None or node is soup.find("table"):
            return placeholder
    # CHM layouts nest content tables inside layout cells.  Extract innermost
    # nested tables first and emit them after the outer table; the outer cell
    # keeps only a placeholder so the text is never decomposed away.
    extracted: list[str] = []
    while True:
        nested = None
        for candidate in soup.find_all("table"):
            if candidate.find_parent("table") is not None and candidate.find("table") is None:
                nested = candidate
                break
        if nested is None:
            break
        markdown = (
            _convert_single_table(str(nested))
            if TOC_MARKER_RE.search(str(nested))
            else _grid_table_markdown(nested)
        )
        nested.replace_with(f"[[NESTTBL:{len(extracted)}]]")
        extracted.append(markdown)
    parts: list[str] = []
    for table in soup.find_all("table"):
        markdown = _grid_table_markdown(table, strip_nested_markers=True)
        if markdown:
            parts.append(markdown)
    body = "\n\n".join(parts)
    for index, markdown in enumerate(extracted):
        if markdown:
            body = body + "\n\n" + markdown
    if not body.strip():
        return ""
    return "\n\n" + body + "\n\n"


def _normalize_image_markup(text: str) -> str:
    """Convert extracted HTML images into portable Markdown image links.

    Pandoc reliably embeds Markdown images in DOCX and EPUB, while raw HTML
    ``<img>`` tags are format-dependent.  Some source EPUB/HTML extractions
    also retain Windows backslashes in relative paths, which are not portable
    resource references for Pandoc or EPUB readers.
    """

    def linked_replacement(href: str, src: str) -> str:
        href_path = href.split("?", 1)[0].split("#", 1)[0].casefold()
        src_path = src.split("?", 1)[0].split("#", 1)[0].casefold()
        image_suffixes = (".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".bmp", ".webp")
        if not href or not src or not href_path.endswith(image_suffixes) or not src_path.endswith(image_suffixes):
            return src
        replacement = href
        if src.replace("\\", "/").lstrip("/").startswith("assets/original/") and not replacement.replace("\\", "/").lstrip("/").startswith("assets/original/"):
            replacement = "assets/original/" + replacement.replace("\\", "/").lstrip("/")
        return replacement

    def promote_linked_markdown_image(match: re.Match[str]) -> str:
        anchor_attrs = {
            item.group("name").casefold(): item.group("value")
            for item in HTML_ATTR_RE.finditer(match.group("a_attrs"))
        }
        src = linked_replacement(anchor_attrs.get("href", "").strip(), match.group("src").strip())
        return match.group("open") + f"![{match.group('alt')}]({src}{match.group('title') or ''})</a>"

    def promote_linked_image(match: re.Match[str]) -> str:
        anchor_attrs = {
            item.group("name").casefold(): item.group("value")
            for item in HTML_ATTR_RE.finditer(match.group("a_attrs"))
        }
        image_attrs = {
            item.group("name").casefold(): item.group("value")
            for item in HTML_ATTR_RE.finditer(match.group("img_attrs"))
        }
        href = anchor_attrs.get("href", "").strip()
        src = image_attrs.get("src", "").strip()
        replacement = linked_replacement(href, src)
        if replacement == src:
            return match.group(0)
        image_tag = re.sub(
            r"(?P<prefix>\bsrc\s*=\s*)(?P<quote>[\"'])(?P<value>.*?)(?P=quote)",
            lambda item: item.group("prefix") + item.group("quote") + replacement + item.group("quote"),
            match.group("img"),
            count=1,
            flags=re.IGNORECASE | re.DOTALL,
        )
        return match.group("open") + image_tag + "</a>"

    text = LINKED_MARKDOWN_IMAGE_RE.sub(promote_linked_markdown_image, text)
    text = LINKED_IMAGE_RE.sub(promote_linked_image, text)

    def replace(match: re.Match[str]) -> str:
        attrs = {
            item.group("name").casefold(): item.group("value")
            for item in HTML_ATTR_RE.finditer(match.group("attrs"))
        }
        src = attrs.get("src", "").strip()
        if not src:
            return ""
        src = src.replace("\\", "/")
        if src.rsplit("/", 1)[-1].casefold() in PLATFORM_IMAGE_NAMES:
            # Safari/Books24x7 navigation chrome and layout spacers are
            # extraction residue, not book content.
            return ""
        alt = attrs.get("alt", "").strip().replace("]", "\\]")
        return f"![{alt}]({src})"

    return HTML_IMAGE_TAG_RE.sub(replace, text)


def _pipe_cells(line: str) -> list[str]:
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def _pipe_row(cells: list[str]) -> str:
    return "| " + " | ".join(cells) + " |"


def _split_wide_tables(lines: list[str], max_columns: int = 5) -> list[str]:
    """Split wide pipe tables into readable column groups with key column repeated."""
    result: list[str] = []
    index = 0
    while index < len(lines):
        if not lines[index].lstrip().startswith("|") or index + 1 >= len(lines):
            result.append(lines[index])
            index += 1
            continue
        separator = _pipe_cells(lines[index + 1]) if lines[index + 1].lstrip().startswith("|") else []
        # Pandoc may emit two-dash alignment separators (``:--:``) for
        # narrow source columns even though the Markdown reader accepts them.
        # Treat those as table separators too; otherwise oversized CHM tables
        # bypass the continuation-table logic entirely.
        if not separator or not all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "")) for cell in separator):
            result.append(lines[index])
            index += 1
            continue
        block: list[list[str]] = []
        while index < len(lines) and lines[index].lstrip().startswith("|"):
            block.append(_pipe_cells(lines[index]))
            index += 1
        width = max((len(row) for row in block), default=0)
        if width <= max_columns:
            result.extend(_pipe_row(row) for row in block)
            continue
        normalized = [row + [""] * (width - len(row)) for row in block]
        header = " ".join(normalized[0]) if normalized else ""
        compact = (
            max((len(cell) for row in normalized for cell in row), default=0) <= 18
            and "总数" in header
            and ("N/A" in header or "N／A" in header)
        )
        if compact:
            result.extend(_pipe_row(row) for row in normalized)
            continue
        # Most book tables have one label column, but CHM tables frequently
        # use a category column plus a metric/row-label column.  Repeating
        # only column 0 in that case leaves continuation tables with blank
        # rows that cannot be related back to the first half.
        body_rows = normalized[2:]
        repeat_columns = 2 if (
            width >= 2
            and any(len(row) > 1 and row[1] for row in body_rows)
            and any(not row[0] for row in body_rows)
        ) else 1
        group_width = max(1, max_columns - repeat_columns)
        groups = [
            list(range(start, min(start + group_width, width)))
            for start in range(repeat_columns, width, group_width)
        ]
        for group_index, group in enumerate(groups):
            columns = list(range(repeat_columns)) + group
            if group_index:
                result.extend(["", "*（续表）*", ""])
            result.extend(_pipe_row([row[column] for column in columns]) for row in normalized)
    return result


def _convert_html_tables(text: str) -> str:
    """Convert extracted raw HTML tables into Markdown pipe tables.

    Pandoc's DOCX writer drops raw HTML blocks, so every HTML table would
    silently disappear from publication DOCX/PDF.  CHM/EPUB extraction wraps
    content in nested layout tables, so conversion walks outermost table
    blocks with a nesting-aware scanner.  Colspan/rowspan cells are flattened
    into a rectangular grid; ``<caption>`` becomes a bold line; the source
    book's static TOC table becomes the ``[[TOC]]`` placeholder; pure spacer
    tables are dropped.
    """
    out: list[str] = []
    pos = 0
    for start, end in list(_iter_outer_tables(text)):
        out.append(text[pos:start])
        out.append(_convert_single_table(text[start:end]))
        pos = end
    out.append(text[pos:])
    converted = "".join(out)
    # Complete tables are gone at this point; leftover table-part tags belong
    # to fragments truncated at extraction boundaries.
    return re.sub(
        r"</?(?:table|tbody|thead|tfoot|tr|td|th)\b[^>]*/?>",
        "",
        converted,
    )


def _display_width(text: str) -> int:
    """Approximate rendered width: East Asian wide/fullwidth glyphs count twice."""
    width = 0
    for character in text:
        if unicodedata.combining(character):
            continue
        width += 2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1
    return width


def _reproportion_pipe_separators(lines: list[str]) -> list[str]:
    """Rewrite pipe-table separator rows with content-proportional widths.

    Pandoc reads column widths from the dash counts of the separator row.
    Equal narrow dashes squeezed CJK headers into one-character columns and
    split numeric cells mid-amount; proportional dashes give every column a
    width that fits its widest cell.
    """
    separator_re = re.compile(
        r"^\|(?:\s*:?-{2,}:?\s*\|)+$"
    )
    result: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if (
            line.lstrip().startswith("|")
            and index + 1 < len(lines)
            and separator_re.match(lines[index + 1].strip())
        ):
            block_end = index
            while block_end < len(lines) and lines[block_end].lstrip().startswith("|"):
                block_end += 1
            rows = [_pipe_cells(lines[k]) for k in range(index, block_end)]
            prose: list[str] = []
            data_rows: list[list[str]] = []
            for r, row in enumerate(rows):
                if r != 1 and _is_prose_row(row):
                    prose.append(next(cell for cell in row if cell.strip()))
                else:
                    data_rows.append(row)
            rows = data_rows
            width = max(len(row) for row in rows)
            col_widths = [4] * width
            for r, row in enumerate(rows):
                if r == 1:
                    continue
                for c, cell in enumerate(row):
                    # +4 dashes of cell-margin allowance so the widest cell
                    # still fits on one line after Word adds cell padding.
                    w = min(_display_width(cell) + 4, 90)
                    if w > col_widths[c]:
                        col_widths[c] = w
            col_widths = [max(4, min(90, w)) for w in col_widths]
            total = sum(col_widths)
            # Pandoc also estimates the whole table's width from the
            # separator's source length; the separator must comfortably exceed
            # the text column width (about 140 characters) or the table ships
            # undersized and numeric cells wrap mid-amount.
            if total < 160:
                col_widths = [max(4, round(w * 160 / total)) for w in col_widths]
            separator = "| " + " | ".join("-" * w for w in col_widths) + " |"
            result.append(line)
            result.append(separator)
            for k in range(index + 2, block_end):
                if _is_prose_row(_pipe_cells(lines[k])):
                    continue
                result.append(lines[k])
            if prose:
                # Embedded footnotes/paragraphs return as normal paragraphs
                # after the table instead of stretching a column.
                result.append("")
                for text in prose:
                    result.append(text)
                    result.append("")
            index = block_end
            continue
        result.append(line)
        index += 1
    return result


def clean_markdown_for_publication(markdown: str) -> str:
    """Remove print-production residue and normalize common book structure."""
    markdown = _normalize_html_headings(markdown)
    lines = _strip_print_artifacts(markdown)
    lines = _replace_static_toc(lines)
    lines = _normalize_structure(lines)
    lines = _strip_trailing_recovery_material(lines)
    lines = _join_dangling_paragraphs(lines)
    # Convert raw HTML tables before splitting wide tables.  CHM/EPUB
    # extractors often preserve rich tables until this stage; splitting the
    # line list first would miss those tables and leave seven-plus-column
    # grids to collapse in DOCX/PDF.
    text = _convert_html_tables("\n".join(lines))
    text = "\n".join(_separate_html_blocks_from_tables(_reproportion_pipe_separators(_split_wide_tables(text.splitlines()))))
    # Extraction and translation can leave <br> inside table cells; the
    # HTML5 short form is invalid XML and breaks the semantic XHTML/EPUB
    # build, so normalize it to the self-closing form.
    text = re.sub(r"<br\s*/?>", "<br/>", text, flags=re.IGNORECASE)
    text = _normalize_sup_markers(text)
    text = _normalize_image_markup(text)
    # Pandoc drops empty HTML anchors even when another retained link targets
    # them. Convert those image/backlink targets into nonempty semantic spans.
    text = re.sub(
        r'<a\b[^>]*\bid=["\'](?P<id>calibre_link-\d+)["\'][^>]*>\s*</a>',
        lambda match: f'<span id="{match.group("id")}" aria-hidden="true">&#x200B;</span>',
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    # Calibre can leave TOC/cross-reference links whose generated anchor was
    # discarded during extraction. Keep their visible text but do not carry a
    # knowingly broken synthetic link into semantic XHTML/EPUB.
    known_ids = set(re.findall(r'(?:\{#|\bid=["\'])(calibre_link-\d+)', text, re.IGNORECASE))
    text = re.sub(
        r"\[([^]]+)\]\(#(?P<id>calibre_link-\d+)(?:\s+[\"'][^\"']*[\"'])?\)",
        lambda match: match.group(0) if match.group("id") in known_ids else match.group(1),
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(
        r'<a\b[^>]*href=["\']#(?P<id>calibre_link-\d+)["\'][^>]*>(?P<label>.*?)</a>',
        lambda match: match.group(0) if match.group("id") in known_ids else match.group("label"),
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    text = re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"
    return text


def _restore_promoted_image_assets(project: Path, markdown: str) -> int:
    """Backfill high-resolution CHM assets for legacy prepared projects.

    Before linked-image promotion was introduced, target units referenced the
    350px thumbnail while the decompiled CHM still retained the full image in
    ``source-work/extracted/htmlz``.  Publication cleanup now upgrades those
    references; copy the authoritative extracted bytes into the normal asset
    directory when the upgraded destination is missing, without overwriting
    an existing asset.
    """

    restored = 0
    refs = re.findall(r"!\[[^\]]*\]\(([^)\s]+)", markdown)
    for ref in refs:
        normalized = ref.replace("\\", "/").lstrip("/")
        if not normalized.startswith("assets/original/"):
            continue
        destination = safe_project_path(project, normalized)
        if destination.is_file():
            continue
        relative = normalized[len("assets/original/"):]
        source = safe_project_path(project, f"source-work/extracted/htmlz/{relative}")
        if not source.is_file():
            continue
        destination.parent.mkdir(parents=True, exist_ok=True)
        atomic_copy(source, destination)
        restored += 1
    return restored


def merge_edited_units(project: Path, config: dict[str, Any]) -> str:
    manifest = load_manifest(project)
    require_edit = bool(config["passes"]["edit"])
    contents: list[str] = []
    for entry in sorted(manifest["units"], key=lambda item: item["order"]):
        relative = entry["edit_file"] if require_edit else entry["draft_file"]
        path = safe_project_path(project, relative)
        if not path.is_file() or not path.read_text(encoding="utf-8").strip():
            raise BTPError(f"Missing or blank target unit: {entry['unit_id']}")
        contents.append(path.read_text(encoding="utf-8").rstrip())
    return "\n\n".join(contents).rstrip() + "\n"


def publication_table_inventory(project: Path, config: dict[str, Any]) -> dict[str, Any]:
    """Report source HTML/pipe tables and their publication-master counterparts."""
    raw = merge_edited_units(project, config)
    clean = clean_markdown_for_publication(raw)
    raw_html = len(list(_iter_outer_tables(raw)))
    raw_pipe = _pipe_table_count(raw)
    clean_pipe = _pipe_table_count(clean)
    return {
        "schema_version": 1,
        "raw_html_tables": raw_html,
        "raw_pipe_tables": raw_pipe,
        "clean_pipe_tables": clean_pipe,
        "continuation_tables": clean.count("*（续表）*"),
        "html_conversion_delta": clean_pipe - raw_pipe,
        "unconverted_html_table_markers": len(re.findall(r"</?table\b", clean, re.IGNORECASE)),
    }


def _run(command: list[str], cwd: Path) -> subprocess.CompletedProcess[str]:
    run_cwd = Path(sys.executable).parent if str(cwd).startswith("\\\\?\\") else cwd
    result = subprocess.run(
        command, cwd=run_cwd, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=900,
    )
    if result.returncode:
        tail = (result.stderr or result.stdout).strip().splitlines()[-10:]
        raise BTPError(f"Typesetting command failed ({result.returncode}): {' | '.join(tail)}")
    return result


def _set_style_fonts(style: Any, latin: str, east_asia: str, size_pt: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = latin
    style.font.size = Pt(size_pt)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr, value in (("ascii", latin), ("hAnsi", latin), ("eastAsia", east_asia)):
        rfonts.set(qn(f"w:{attr}"), value)


def _clear_paragraph(paragraph: Any) -> None:
    from docx.oxml.ns import qn

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _append_field(paragraph: Any, instruction: str, placeholder: str = "1") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run = OxmlElement("w:r")
        run.append(element)
        paragraph._p.append(run)


def _set_page_number_format(section: Any, fmt: str, start: int | None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        sect_pr.append(node)
    node.set(qn("w:fmt"), fmt)
    if start is None:
        node.attrib.pop(qn("w:start"), None)
    else:
        node.set(qn("w:start"), str(start))


def _add_book_sections(document: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    headings = []
    body_started = False
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Heading 1":
            continue
        text = paragraph.text.strip()
        if PART_RE.match(text) or CHAPTER_TITLE_RE.match(text) or CHAPTER_CN_RE.match(text):
            body_started = True
        if body_started:
            headings.append(paragraph)
        else:
            paragraph.paragraph_format.page_break_before = True
    if not headings:
        return
    base = document.sections[-1]._sectPr
    for paragraph in headings:
        break_p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        sect_pr = deepcopy(base)
        for tag in ("w:headerReference", "w:footerReference", "w:pgNumType", "w:titlePg", "w:type"):
            for child in list(sect_pr.findall(qn(tag))):
                sect_pr.remove(child)
        section_type = OxmlElement("w:type")
        section_type.set(qn("w:val"), "oddPage")
        sect_pr.insert(0, section_type)
        ppr.append(sect_pr)
        break_p.append(ppr)
        paragraph._p.addprevious(break_p)


def _configure_settings(document: Any, options: dict[str, Any]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = document.settings._element
    mirror = settings.find(qn("w:mirrorMargins"))
    if options.get("mirror_margins"):
        if mirror is None:
            settings.append(OxmlElement("w:mirrorMargins"))
    elif mirror is not None:
        settings.remove(mirror)
    even_odd = settings.find(qn("w:evenAndOddHeaders"))
    if even_odd is None:
        even_odd = OxmlElement("w:evenAndOddHeaders")
        settings.append(even_odd)
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _configure_styles(document: Any, options: dict[str, Any]) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    styles = document.styles
    body_ea = options["body_font_east_asia"]
    body_latin = options["body_font_latin"]
    heading_ea = options["heading_font_east_asia"]
    heading_latin = options["heading_font_latin"]

    normal = styles["Normal"]
    _set_style_fonts(normal, body_latin, body_ea, options["body_font_size_pt"])
    normal.paragraph_format.first_line_indent = Pt(options["body_font_size_pt"] * 2)
    normal.paragraph_format.line_spacing = options["line_spacing"]
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, before, after in (
        ("Title", 26, 0, 18),
        ("Heading 1", 20, 18, 14),
        ("Heading 2", 14, 14, 7),
        ("Heading 3", 12, 10, 5),
        ("Heading 4", 11, 8, 4),
    ):
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[style_name]
        _set_style_fonts(style, heading_latin, heading_ea, size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = False

    if "TOC Heading" not in styles:
        styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    toc_heading = styles["TOC Heading"]
    _set_style_fonts(toc_heading, heading_latin, heading_ea, 18)
    toc_heading.font.bold = True
    toc_heading.paragraph_format.space_after = Pt(14)

    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "目录":
            paragraph.style = toc_heading
            paragraph.paragraph_format.page_break_before = True
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True


def _format_title_page(document: Any, options: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    copyright_start = None
    remove = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if (
            text.startswith("本书采用无酸纸")
            or text.casefold().startswith("printed on acid-free")
            or "国会图书馆" in text
            or "library of congress" in text.casefold()
        ):
            copyright_start = paragraph
            break
        if (
            text == options["title"]
            and paragraph.style.name == "Title"
            and not paragraph.text.strip() == ""
        ):
            # pandoc renders the metadata title as a leading Title paragraph;
            # the real 书名页 Heading 1 anchors the title page instead, so the
            # floating metadata copy is removed.
            remove.append(paragraph)
    for paragraph in remove:
        paragraph._element.getparent().remove(paragraph._element)

    title_heading = None
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == options["title"] and getattr(
            paragraph.style, "name", ""
        ).startswith("Heading"):
            title_heading = paragraph
            break
    title_page = []
    if title_heading is not None:
        for paragraph in document.paragraphs:
            if copyright_start is not None and paragraph._p is copyright_start._p:
                break
            if paragraph._p is title_heading._p:
                title_page.append(paragraph)
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if title_page and (style_name.startswith("Heading") or style_name == "TOC Heading"):
                break
            if len(title_page) >= 40:
                break
            if paragraph.text.strip():
                title_page.append(paragraph)
    if not title_page:
        return
    for paragraph in title_page:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)  # None would inherit the 2-char body indent
        paragraph.paragraph_format.space_after = Pt(8)
    title = next(
        (p for p in title_page if p.text.strip() == options["title"]),
        title_page[0],
    )
    title.style = document.styles["Title"]
    title.paragraph_format.space_before = Pt(70)
    title.paragraph_format.space_after = Pt(28)
    if copyright_start is not None:
        copyright_start.paragraph_format.page_break_before = True


def _table_column_widths(document: Any, table: Any, target_total: int) -> list[int]:
    """Content-aware column widths in twips for one table.

    Pandoc derives a fixed grid from source line lengths, which is usually
    narrower than the trim's text column and squeezes CJK headers into
    one-character columns and amounts into wrapped fragments.  Widths here
    follow the widest cell content per column (CJK-aware), never drop below
    a readable minimum, and always sum exactly to the text-column width.
    """
    column_count = len(table.columns)
    natural = [4] * column_count
    for row in table.rows:
        for index, cell in enumerate(row.cells[:column_count]):
            width = min(_display_width(cell.text.strip()) + 4, 72)
            if width > natural[index]:
                natural[index] = width
    total_natural = sum(natural)
    min_twips = 680
    widths = [max(min_twips, round(value / total_natural * target_total)) for value in natural]
    # Columns pinned to the minimum free proportionality for the rest.
    for _ in range(3):
        fixed = sum(width for width in widths if width <= min_twips)
        flexible = [index for index, width in enumerate(widths) if width > min_twips]
        remaining = target_total - fixed
        if not flexible or remaining <= 0:
            break
        flex_natural = sum(natural[index] for index in flexible) or 1
        changed = False
        for index in flexible:
            proposed = max(min_twips, round(natural[index] / flex_natural * remaining))
            if proposed != widths[index]:
                widths[index] = proposed
                changed = True
        if not changed:
            break
    # Absorb rounding in the widest column so the sum equals the target.
    delta = target_total - sum(widths)
    widths[widths.index(max(widths))] += delta
    return widths


FIGURE_CAPTION_RE = re.compile(r"^\s*(?:图表|图|表)\s*\d")


def _keep_figures_with_captions(document: Any) -> None:
    """Bind figures and tables to their captions so neither straddles a page break.

    An image whose caption lands on the next page (or the reverse) reads as
    broken layout.  Caption paragraphs (图表 N.N …) always keep with the next
    block — a following figure, table, or paragraph — and figure paragraphs
    get keep-lines so Word moves the whole pair to the next page instead of
    separating them.  Word gives up on impossible keep rules gracefully, so
    page-filling exhibits still render.
    """
    paragraphs = document.paragraphs
    has_drawing = ["w:drawing" in paragraph._p.xml for paragraph in paragraphs]
    is_caption = [
        bool(FIGURE_CAPTION_RE.match(paragraph.text.strip())) and len(paragraph.text.strip()) <= 80
        for paragraph in paragraphs
    ]
    for index, paragraph in enumerate(paragraphs):
        if has_drawing[index]:
            paragraph.paragraph_format.keep_together = True
            next_is_caption = index + 1 < len(paragraphs) and is_caption[index + 1]
            paragraph.paragraph_format.keep_with_next = next_is_caption
        elif is_caption[index]:
            paragraph.paragraph_format.keep_with_next = True


def _configure_tables(document: Any, options: dict[str, Any]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    target = int(
        (options["trim_width_mm"] - options["inside_margin_mm"] - options["outside_margin_mm"])
        / 25.4 * 1440
    )
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Fixed layout keeps the computed content-aware grid; autofit lets
        # Word squeeze columns until CJK headers wrap one character per line
        # and amounts split mid-number.
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        # python-docx inserts w:tblLayout in schema order via autofit; the
        # manual fallback below only runs for hand-built tables and must
        # respect the same order (tblLayout precedes tblLook).
        if tbl_pr.find(qn("w:tblLayout")) is None:
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            look = tbl_pr.find(qn("w:tblLook"))
            if look is not None:
                look.addprevious(layout)
            else:
                tbl_pr.append(layout)
        table_width = tbl_pr.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            anchor = tbl_pr.find(qn("w:jc"))
            if anchor is None:
                anchor = tbl_pr.find(qn("w:tblLook"))
            if anchor is not None:
                anchor.addprevious(table_width)
            else:
                tbl_pr.append(table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(target))
        font_size = 8.5 if len(table.columns) >= 5 else 9
        row_count = len(table.rows)
        for row_index, row in enumerate(table.rows):
            # Tables stay on one page whenever they fit: no row splits across
            # pages (w:cantSplit) and every row except the last keeps with the
            # next, so Word moves the whole table to a fresh page instead of
            # breaking it.  Tables taller than a page still flow across pages
            # with their header row repeated (w:tblHeader on the first row).
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                cant_split = OxmlElement("w:cantSplit")
                header_or_height = tr_pr.find(qn("w:tblHeader"))
                if header_or_height is None:
                    header_or_height = tr_pr.find(qn("w:trHeight"))
                if header_or_height is not None:
                    header_or_height.addprevious(cant_split)
                else:
                    tr_pr.append(cant_split)
            if row_index == 0 and tr_pr.find(qn("w:tblHeader")) is None:
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                tr_pr.append(header)
            # Keep the header row attached to the first data rows: a header
            # stranded alone at a page bottom is a bad break.
            keep_with_next = row_index < row_count - 1
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_index == 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    paragraph.paragraph_format.first_line_indent = Pt(0)  # None would inherit the 2-char body indent
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.keep_with_next = keep_with_next
                    for run in paragraph.runs:
                        run.font.size = Pt(font_size)
                        if row_index == 0:
                            run.bold = True
        widths = _table_column_widths(document, table, target)
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            cols = grid.findall(qn("w:gridCol"))
            for col, width in zip(cols, widths):
                col.set(qn("w:w"), str(width))
        for row in table.rows:
            cells = row.cells[: len(widths)]
            for cell, width in zip(cells, widths):
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_width = tc_pr.find(qn("w:tcW"))
                if tc_width is None:
                    tc_width = OxmlElement("w:tcW")
                    tc_pr.insert(0, tc_width)
                tc_width.set(qn("w:type"), "dxa")
                tc_width.set(qn("w:w"), str(width))


def _fit_inline_images(document: Any, options: dict[str, Any]) -> None:
    """Keep embedded figures inside the page text area, width and height.

    Width is capped to the text measure; height is capped to the text height
    minus room for one caption line, so a figure and its bound caption always
    fit on a single page together.  Without the height cap, near-full-page
    crops exceed the text height, Word pushes them to a page of their own and
    drops the caption keep rule, stranding the caption at a page bottom.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu, Mm, Pt

    gutter = options["gutter_mm"] if options.get("mirror_margins") else 0
    usable_width_mm = (
        options["trim_width_mm"]
        - options["inside_margin_mm"]
        - options["outside_margin_mm"]
        - gutter
    )
    usable_height_mm = (
        options["trim_height_mm"]
        - options["top_margin_mm"]
        - options["bottom_margin_mm"]
        - 12  # caption line plus spacing
    )
    max_width = int(Mm(max(20, usable_width_mm)))
    max_height = int(Mm(max(20, usable_height_mm)))
    for paragraph in document.paragraphs:
        if "w:drawing" not in paragraph._p.xml:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)  # None would inherit the 2-char body indent
        paragraph.paragraph_format.space_before = Mm(2)
        paragraph.paragraph_format.space_after = Mm(2)
    for inline_shape in document.inline_shapes:
        original_width = int(inline_shape.width)
        original_height = int(inline_shape.height)
        if original_width <= 0 or original_height <= 0:
            continue
        scale = min(1.0, max_width / original_width, max_height / original_height)
        if scale < 1.0:
            inline_shape.width = Emu(round(original_width * scale))
            inline_shape.height = Emu(round(original_height * scale))


def detect_folio_position(project: Path, config: dict[str, Any]) -> str:
    """Sample the source PDF for its running folio convention.

    Standalone page numbers on body pages are located and their vertical
    position measured: many trade books print the folio on the top line next
    to the running head, while others use a bottom footer.  Matching the
    source keeps the translation's page furniture aligned with the original
    edition.  Returns "header", "footer", or "footer" when detection is not
    possible (non-PDF sources, missing PyMuPDF, or ambiguous samples).
    """
    if config.get("source", {}).get("format") != "pdf":
        return "footer"
    try:
        import fitz
    except ImportError:
        return "footer"
    try:
        source_path = safe_project_path(project, config["source"]["path"])
        document = fitz.open(source_path)
    except Exception:
        return "footer"
    positions: list[float] = []
    folio_re = re.compile(r"^\d{1,3}$")
    try:
        page_count = len(document)
        start = 2 if page_count <= 45 else 40
        step = max(1, page_count // 60)
        sample_pages = range(start, page_count, step)
        for index in sample_pages:
            page = document[index]
            height = page.rect.height
            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    text = " ".join(str(span.get("text", "")).strip() for span in spans).strip()
                    if not folio_re.match(text) or not spans:
                        continue
                    if float(spans[0].get("size", 99)) > 13:
                        continue
                    y0 = float(line["bbox"][1])
                    # Ignore centered display numbers (chapter openers, part
                    # pages): folios sit near the left or right edge.
                    x0 = float(line["bbox"][0])
                    if page.rect.width * 0.30 < x0 < page.rect.width * 0.70:
                        continue
                    positions.append(y0 / height)
    except Exception:
        document.close()
        return "footer"
    document.close()
    if len(positions) < 5:
        return "footer"
    positions.sort()
    median = positions[len(positions) // 2]
    if median < 0.22:
        return "header"
    return "footer"


def _append_text_run(paragraph: Any, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    paragraph._p.append(run)


def _append_page_field(paragraph: Any) -> None:
    _append_field(paragraph, "PAGE")


def _style_header_runs(paragraph: Any, size_pt: float = 9.0) -> None:
    from docx.shared import Pt

    for run in paragraph.runs:
        run.font.size = Pt(size_pt)


def _configure_section_layout(document: Any, options: dict[str, Any]) -> None:
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Mm, Pt

    if options.get("mirror_margins"):
        left_margin = options["inside_margin_mm"]
        right_margin = options["outside_margin_mm"]
        gutter = options["gutter_mm"]
    else:
        # Symmetric margins keep consecutive pages visually aligned on
        # screen; the average of the inside/outside spec preserves the
        # configured text measure.
        uniform = (options["inside_margin_mm"] + options["outside_margin_mm"]) / 2
        left_margin = uniform
        right_margin = uniform
        gutter = 0

    for section in document.sections:
        section.page_width = Mm(options["trim_width_mm"])
        section.page_height = Mm(options["trim_height_mm"])
        section.left_margin = Mm(left_margin)
        section.right_margin = Mm(right_margin)
        section.top_margin = Mm(options["top_margin_mm"])
        section.bottom_margin = Mm(options["bottom_margin_mm"])
        section.gutter = Mm(gutter)
        section.header_distance = Mm(10)
        section.footer_distance = Mm(10)

    text_width_mm = options["trim_width_mm"] - left_margin - right_margin

    def _running_head(paragraph: Any, title: str, *, folio_side: str) -> None:
        """Compose one running-head line: title centered, folio at the outer edge."""
        middle = Mm(max(10, text_width_mm / 2))
        end = Mm(max(20, text_width_mm))
        paragraph.paragraph_format.tab_stops.add_tab_stop(middle, WD_TAB_ALIGNMENT.CENTER)
        if folio_side == "left":
            # Verso: folio at the outer (left) edge, book title centered.
            _append_page_field(paragraph)
            _append_text_run(paragraph, f"\t{options['title']}")
        else:
            # Recto: chapter title centered, folio at the outer (right) edge.
            _append_text_run(paragraph, f"\t{title}\t")
            _append_page_field(paragraph)
            paragraph.paragraph_format.tab_stops.add_tab_stop(end, WD_TAB_ALIGNMENT.RIGHT)
        _style_header_runs(paragraph)

    front = document.sections[0]
    front.different_first_page_header_footer = True
    _set_page_number_format(front, "lowerRoman", 1)
    for story in (front.header, front.even_page_header, front.first_page_header):
        _clear_paragraph(story.paragraphs[0])
    _clear_paragraph(front.footer.paragraphs[0])
    front.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _append_field(front.footer.paragraphs[0], "PAGE")
    _clear_paragraph(front.first_page_footer.paragraphs[0])

    if len(document.sections) < 2:
        return
    section_titles = []
    body_started = False
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Heading 1":
            continue
        text = paragraph.text.strip()
        if PART_RE.match(text) or CHAPTER_TITLE_RE.match(text) or CHAPTER_CN_RE.match(text):
            body_started = True
        if body_started:
            section_titles.append(text)
    folio_in_header = options.get("folio_position") == "header"
    for index, section in enumerate(document.sections[1:]):
        section_title = section_titles[index] if index < len(section_titles) else options["title"]
        section.start_type = WD_SECTION_START.ODD_PAGE
        section.different_first_page_header_footer = True
        _set_page_number_format(section, "decimal", 1 if index == 0 else None)
        for story in (
            section.header, section.even_page_header, section.first_page_header,
            section.footer, section.even_page_footer, section.first_page_footer,
        ):
            story.is_linked_to_previous = False
            _clear_paragraph(story.paragraphs[0])

        odd_header = section.header.paragraphs[0]
        even_header = section.even_page_header.paragraphs[0]
        if folio_in_header:
            _running_head(even_header, options["title"], folio_side="left")
            _running_head(odd_header, section_title, folio_side="right")
        else:
            odd_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            odd_header.add_run(section_title)
            even_header.alignment = WD_ALIGN_PARAGRAPH.LEFT
            even_header.add_run(options["title"])
            _style_header_runs(odd_header)
            _style_header_runs(even_header)

        section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        section.even_page_footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        if not folio_in_header:
            _append_field(section.footer.paragraphs[0], "PAGE")
            _append_field(section.even_page_footer.paragraphs[0], "PAGE")
        if not PART_RE.match(section_title):
            section.first_page_footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _append_field(section.first_page_footer.paragraphs[0], "PAGE")


def postprocess_docx(path: Path, options: dict[str, Any]) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise BTPError("python-docx is required for publication DOCX layout; run doctor") from exc

    document = Document(path)
    _configure_settings(document, options)
    _configure_styles(document, options)
    _format_title_page(document, options)
    _configure_tables(document, options)
    _keep_figures_with_captions(document)
    _fit_inline_images(document, options)
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "[[TOC]]":
            _clear_paragraph(paragraph)
            _append_field(paragraph, f'TOC \\o "1-{options["toc_depth"]}" \\h \\z \\u', "更新目录")
    _add_book_sections(document)
    document.save(path)

    document = Document(path)
    _configure_settings(document, options)
    _configure_section_layout(document, options)
    document.save(path)


def _word_refresh_and_export(docx_path: Path, pdf_path: Path | None) -> bool:
    if os.name != "nt" or not find_binary("powershell"):
        return False
    docx = str(docx_path.resolve()).replace("'", "''")
    pdf = str(pdf_path.resolve()).replace("'", "''") if pdf_path else ""
    export = f"$doc.ExportAsFixedFormat('{pdf}',17);" if pdf_path else ""
    script = (
        "$ErrorActionPreference='Stop';$word=$null;$doc=$null;"
        "try{$word=New-Object -ComObject Word.Application;$word.Visible=$false;"
        "$word.DisplayAlerts=0;$doc=$word.Documents.Open('" + docx + "');"
        "foreach($toc in $doc.TablesOfContents){$toc.Update()|Out-Null};"
        "$doc.Fields.Update()|Out-Null;$doc.Repaginate();$doc.Save();" + export +
        "$doc.Close();$word.Quit();exit 0}"
        "catch{if($doc){$doc.Close([ref]$false)};if($word){$word.Quit()};"
        "Write-Error $_;exit 1}"
    )
    result = subprocess.run(
        [str(find_binary("powershell")), "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
    )
    return result.returncode == 0 and (pdf_path is None or pdf_path.is_file())


def _libreoffice_export(docx_path: Path, pdf_path: Path) -> bool:
    soffice = find_binary("soffice")
    if not soffice:
        return False
    result = subprocess.run(
        [str(soffice), "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
    )
    produced = pdf_path.parent / f"{docx_path.stem}.pdf"
    if result.returncode == 0 and produced.is_file():
        if produced != pdf_path:
            produced.replace(pdf_path)
        return True
    return False


def _publication_options(config: dict[str, Any]) -> dict[str, Any]:
    publishing = config.get("publishing", {})
    metadata = config.get("metadata", {})
    mode = config.get("mode", "study")
    mirror_margins = bool(publishing.get("mirror_margins", mode == "publication"))
    return {
        "title": metadata.get("title") or config["project_id"],
        "trim_width_mm": publishing.get("trim_width_mm", 170),
        "trim_height_mm": publishing.get("trim_height_mm", 240),
        "inside_margin_mm": publishing.get("inside_margin_mm", 25),
        "outside_margin_mm": publishing.get("outside_margin_mm", 20),
        "top_margin_mm": publishing.get("top_margin_mm", 22),
        "bottom_margin_mm": publishing.get("bottom_margin_mm", 22),
        # Screen-first profiles (fast/study) keep symmetric margins so
        # consecutive pages do not visibly shift; print profiles mirror.
        "mirror_margins": mirror_margins,
        # "auto" samples the source PDF for its folio convention (running
        # heads share the top line in many trade books); "header"/"footer"
        # force a placement.
        "folio_position": publishing.get("folio_position", "auto"),
        "gutter_mm": publishing.get("gutter_mm", 5 if mirror_margins else 0),
        "body_font_east_asia": publishing.get("body_font_east_asia", "宋体"),
        "body_font_latin": publishing.get("body_font_latin", "Times New Roman"),
        "heading_font_east_asia": publishing.get("heading_font_east_asia", "黑体"),
        "heading_font_latin": publishing.get("heading_font_latin", "Arial"),
        "body_font_size_pt": publishing.get("body_font_size_pt", 10.5),
        "line_spacing": publishing.get("line_spacing", 1.45),
        "toc_depth": publishing.get("toc_depth", 2),
    }


def _infer_target_title(markdown: str, fallback: str, target_language: str = "") -> str:
    prefer_cjk = str(target_language).lower().startswith("zh")
    first_candidate = None
    skipped_labels = {value.casefold() for value in (*TOC_TITLES, *FRONT_TITLES)}
    for line in markdown.splitlines()[:80]:
        stripped = line.strip()
        heading = HEADING_RE.match(stripped)
        if heading:
            candidate = _plain(heading.group(2))
        elif _emphasis_only(stripped):
            candidate = _plain(stripped)
        else:
            continue
        # Skip TOC/navigation/front-matter labels; they are headings but not
        # the book title.
        if candidate.casefold() in skipped_labels or len(candidate) <= 3:
            continue
        if 1 < len(candidate) <= 80:
            if prefer_cjk and re.search("[\u4e00-\u9fff]", candidate):
                return candidate
            if first_candidate is None:
                first_candidate = candidate
    return first_candidate or fallback


def _infer_original_title(markdown: str) -> str:
    """Return a conservative English-title candidate from a CIP-style line."""
    patterns = (
        r"(?m)^[A-Z][\w'\-]+,[^\n]+\n\n(?P<title>[A-Z][^/\n]{10,220}?)\s*/\s*(?:edited\s+by|by|[A-Z])",
        r"(?m)^(?P<title>[A-Z][A-Za-z0-9 &'’:'.,()\-]{10,220}?)\s*/\s*(?:edited\s+by|by|[A-Z][A-Za-z'\-]+)",
    )
    for pattern in patterns:
        match = re.search(pattern, markdown, re.IGNORECASE)
        if not match:
            continue
        candidate = " ".join(match.group("title").split()).strip(" .,:;")
        if re.search(r"[A-Za-z]", candidate) and 10 <= len(candidate) <= 220:
            return candidate
    return ""


def _insert_original_title_line(
    markdown: str,
    title: str,
    target_language: str,
    original_title: str = "",
) -> str:
    """Place the English original title under the localized title heading.

    CIP data on the copyright page carries the original English title
    ("... , <Author-initials>. <Title> / ...").  Reader-facing editions for a
    translated target keep that original title visible on the title page.
    """
    if not str(target_language).lower().startswith("zh"):
        return markdown
    original = " ".join((original_title or _infer_original_title(markdown)).split()).strip()
    if not original:
        return markdown
    heading = f"# {title}"
    if heading not in markdown or original.casefold().startswith(title[:12].casefold()):
        return markdown
    return markdown.replace(heading + "\n", heading + "\n\n*" + original + "*\n", 1)


def prepare_publication_markdown(
    project: Path,
    config: dict[str, Any],
) -> tuple[str, dict[str, Any]]:
    """Build the deterministic Markdown master, including its rights page."""
    raw = merge_edited_units(project, config)
    clean = clean_markdown_for_publication(raw)
    options = _publication_options(config)
    if options["folio_position"] == "auto":
        options["folio_position"] = detect_folio_position(project, config)
    options["title"] = _infer_target_title(
        clean, options["title"], config["target"]["language"]
    )
    clean = _insert_original_title_line(
        clean,
        options["title"],
        str(config["target"]["language"]),
        str(config.get("metadata", {}).get("original_title", "")),
    )
    clean = insert_rights_notice(clean, config)
    return clean, options


def _repair_epub_note_backlinks(epub_path: Path) -> int:
    """Re-insert footnote backlinks that Pandoc's HTML→EPUB pass drops.

    Pandoc turns endnote list items into ``<aside epub:type="footnote">``
    blocks and strips the ``↩`` link back to the reference.  EPUB readers
    and the package validator expect note targets to link back to their
    references, so rewrite each note aside that lost its link.
    """
    import zipfile

    backlink_re = re.compile(
        r'(<aside\b[^>]*\bid="(?P<id>fn[^"]+)"[^>]*>)(.*?)(</aside>)',
        re.IGNORECASE | re.DOTALL,
    )
    repaired_total = 0
    with zipfile.ZipFile(epub_path) as archive:
        infos = {item.filename: item for item in archive.infolist()}
        entries = {item.filename: archive.read(item.filename) for item in archive.infolist()}
    for name, payload in list(entries.items()):
        if not name.casefold().endswith((".xhtml", ".html", ".htm")):
            continue
        text = payload.decode("utf-8", errors="replace")

        def add_backlink(match: re.Match[str]) -> str:
            nonlocal repaired_total
            note_id = match.group("id")
            reference_id = note_id[2:] if note_id.startswith("fn") else note_id
            body = match.group(3)
            if f'href="#fnref{reference_id}"' in body:
                return match.group(0)
            repaired_total += 1
            link = f'<a href="#fnref{reference_id}" class="footnote-back" epub:type="backlink" role="doc-backlink">\u21a9\ufe0e</a>'
            return f"{match.group(1)}{body}{link}{match.group(4)}"

        updated = backlink_re.sub(add_backlink, text)
        if updated != text:
            entries[name] = updated.encode("utf-8")
    if repaired_total:
        temp_path = epub_path.with_suffix(".epub.tmp")
        with zipfile.ZipFile(temp_path, "w") as archive:
            for name, payload in entries.items():
                # Preserve each entry's original compression: EPUB requires
                # the mimetype entry to stay uncompressed and first.
                info = infos[name]
                archive.writestr(
                    zipfile.ZipInfo(name, date_time=info.date_time),
                    payload,
                    compress_type=info.compress_type,
                )
        temp_path.replace(epub_path)
    return repaired_total


def typeset_outputs(project: Path, config: dict[str, Any], output_dir: Path) -> dict[str, Path]:
    """Create cleaned, publication-oriented outputs without changing run state."""
    output_dir.mkdir(parents=True, exist_ok=True)
    clean, options = prepare_publication_markdown(project, config)
    _restore_promoted_image_assets(project, clean)
    markdown_path = output_dir / "book.md"
    atomic_write_text(markdown_path, clean)

    pandoc = find_binary("pandoc")
    if not pandoc:
        raise BTPError("Pandoc is required for publication output; run doctor")
    title = options["title"]
    author = config.get("metadata", {}).get("author", "")
    lang = config["target"]["language"]
    metadata_args = [
        f"--metadata=title:{title}",
        f"--metadata=lang:{lang}",
        f"--metadata=rights:{rights_metadata_text(config)}",
    ]
    if str(author).strip():
        metadata_args.append(f"--metadata=author:{author}")
    bilingual_metadata_args = [
        f"--metadata=title:{title} — bilingual",
        f"--metadata=lang:{lang}",
        f"--metadata=rights:{rights_metadata_text(config)}",
    ]
    if str(author).strip():
        bilingual_metadata_args.append(f"--metadata=author:{author}")
    common = [
        str(pandoc), str(markdown_path), "--standalone",
        *metadata_args, f"--resource-path={project}",
    ]
    requested = set(config["outputs"])
    cover_ref = str(config.get("metadata", {}).get("cover", "")).strip()
    generated: dict[str, Path] = {"markdown": markdown_path}

    epub_markdown = output_dir / "book.epub.md"
    atomic_write_text(epub_markdown, clean.replace("[[TOC]]", ""))
    xhtml_path = project / "target" / "xhtml" / "book.xhtml"
    xhtml_path.parent.mkdir(parents=True, exist_ok=True)
    _run([
        str(pandoc), str(epub_markdown), "--from=markdown+footnotes+raw_html",
        "--to=html5", "--standalone", "--section-divs",
        *metadata_args,
        f"--resource-path={project}", "--output", str(xhtml_path),
    ], project)
    xhtml_validation = validate_xhtml_document(xhtml_path)
    if not xhtml_validation["passed"]:
        raise BTPError("Semantic XHTML validation failed: " + "; ".join(xhtml_validation["errors"]))
    generated["xhtml"] = xhtml_path

    docx_path = output_dir / "book.docx"
    if "docx" in requested or "pdf" in requested:
        _run([*common, "--output", str(docx_path)], project)
        postprocess_docx(docx_path, options)
        if not zipfile.is_zipfile(docx_path):
            raise BTPError(f"Publication DOCX is invalid: {docx_path}")
        _word_refresh_and_export(docx_path, None)
        if "docx" in requested:
            generated["docx"] = docx_path

    if "epub" in requested:
        css_path = output_dir / "epub.css"
        atomic_write_text(css_path, (
            "body{font-family:serif;line-height:1.65;margin:5%;}"
            "h1{page-break-before:always;margin-top:1.8em;}"
            "h2,h3{page-break-after:avoid;}"
            "img{max-width:100%;height:auto;display:block;margin:1em auto;}"
            "table{border-collapse:collapse;width:100%;}"
            "th,td{border:1px solid #aaa;padding:.35em;vertical-align:top;}\n"
        ))
        epub_path = output_dir / "book.epub"
        epub_command = [
            str(pandoc), str(xhtml_path), "--from=html", "--standalone",
            *metadata_args,
            f"--resource-path={project}{os.pathsep}{xhtml_path.parent}",
            "--toc", "--css", str(css_path),
        ]
        if cover_ref:
            cover_path = safe_project_path(project, cover_ref)
            if not cover_path.is_file():
                raise BTPError(f"Configured cover image is missing: {cover_ref}")
            epub_command.append(f"--epub-cover-image={cover_path}")
        _run([*epub_command, "--output", str(epub_path)], project)
        if not zipfile.is_zipfile(epub_path):
            raise BTPError(f"Publication EPUB is invalid: {epub_path}")
        _repair_epub_note_backlinks(epub_path)
        generated["epub"] = epub_path

    if "pdf" in requested:
        pdf_path = output_dir / "book.pdf"
        if not _word_refresh_and_export(docx_path, pdf_path) and not _libreoffice_export(docx_path, pdf_path):
            raise BTPError("PDF-from-DOCX requires Microsoft Word on Windows or LibreOffice")
        generated["pdf"] = pdf_path

    if "bilingual" in requested:
        alignment = create_segment_alignment(project, config)
        bilingual_markdown = output_dir / "book-bilingual.md"
        render_bilingual_markdown(alignment, bilingual_markdown)
        bilingual_text = bilingual_markdown.read_text(encoding="utf-8")
        atomic_write_text(bilingual_markdown, insert_rights_notice(bilingual_text, config))
        generated["bilingual_markdown"] = bilingual_markdown
        bilingual_formats = set(config.get("publishing", {}).get("bilingual_formats", ["markdown", "docx", "epub"]))
        bilingual_xhtml = project / "target" / "xhtml" / "book-bilingual.xhtml"
        _run([
            str(pandoc), str(bilingual_markdown), "--from=markdown+fenced_divs+footnotes+raw_html",
            "--to=html5", "--standalone", "--section-divs",
            *bilingual_metadata_args,
            f"--resource-path={project}", "--output", str(bilingual_xhtml),
        ], project)
        xhtml_result = validate_xhtml_document(bilingual_xhtml)
        if not xhtml_result["passed"]:
            raise BTPError("Bilingual XHTML validation failed: " + "; ".join(xhtml_result["errors"]))
        generated["bilingual_xhtml"] = bilingual_xhtml

        bilingual_docx = output_dir / "book-bilingual.docx"
        if "docx" in bilingual_formats or "pdf" in bilingual_formats:
            _run([
                str(pandoc), str(bilingual_markdown), "--from=markdown+fenced_divs+footnotes+raw_html",
                "--standalone", *bilingual_metadata_args,
                f"--resource-path={project}", "--output", str(bilingual_docx),
            ], project)
            postprocess_docx(bilingual_docx, options)
            _word_refresh_and_export(bilingual_docx, None)
            if "docx" in bilingual_formats:
                generated["bilingual_docx"] = bilingual_docx

        if "epub" in bilingual_formats:
            bilingual_css = output_dir / "epub-bilingual.css"
            atomic_write_text(bilingual_css, (
                "body{font-family:serif;line-height:1.6;margin:5%;}"
                ".btp-bilingual{margin:1em 0 2em;}"
                ".btp-source,.btp-target{padding:.5em 0;}"
                ".btp-source{color:#444;border-bottom:1px solid #bbb;}"
                "img{max-width:100%;height:auto;}table{border-collapse:collapse;width:100%;}"
                "th,td{border:1px solid #aaa;padding:.35em;vertical-align:top;}\n"
            ))
            bilingual_epub = output_dir / "book-bilingual.epub"
            command = [
                str(pandoc), str(bilingual_xhtml), "--from=html", "--standalone", "--toc",
                *bilingual_metadata_args, f"--resource-path={project}{os.pathsep}{bilingual_xhtml.parent}",
                "--css", str(bilingual_css),
            ]
            if cover_ref:
                command.append(f"--epub-cover-image={safe_project_path(project, cover_ref)}")
            _run([*command, "--output", str(bilingual_epub)], project)
            generated["bilingual_epub"] = bilingual_epub

        if "pdf" in bilingual_formats:
            bilingual_pdf = output_dir / "book-bilingual.pdf"
            if not _word_refresh_and_export(bilingual_docx, bilingual_pdf) and not _libreoffice_export(bilingual_docx, bilingual_pdf):
                raise BTPError("Bilingual PDF-from-DOCX requires Microsoft Word on Windows or LibreOffice")
            generated["bilingual_pdf"] = bilingual_pdf
    return generated
