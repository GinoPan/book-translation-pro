from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import fitz
from PIL import Image
from docx import Document

from book_translation_pro.figures import (
    _cluster_rects,
    _figure_crop_box,
    _trim_scan_border,
    looks_like_inverted_lineart,
    normalize_image_background,
)
from book_translation_pro.publication import (
    _configure_settings,
    _keep_figures_with_captions,
    _publication_options,
    _table_column_widths,
    _configure_tables,
    detect_folio_position,
)
from book_translation_pro.qa import _number_tokens


class FakePage:
    """Minimal page stub exposing the geometry API the cropper consumes."""

    def __init__(self, rect: fitz.Rect, drawings: list[fitz.Rect], images: list[fitz.Rect] | None = None) -> None:
        self.rect = rect
        self._drawings = [{"rect": rect} for rect in drawings]
        self._images = images or []

    def get_drawings(self) -> list[dict]:
        return self._drawings

    def get_images(self, full: bool = True) -> list[tuple]:
        return [(index,) for index in range(len(self._images))]

    def get_image_rects(self, xref: int) -> list[fitz.Rect]:
        return [self._images[xref]]


class FigureCropTests(unittest.TestCase):
    page_rect = fitz.Rect(0, 0, 433.2, 649.2)

    def test_header_rule_does_not_inflate_crop(self) -> None:
        # Chart cluster in the lower half plus a thin page-header rule.
        chart = [fitz.Rect(77, 358, 374, 491), fitz.Rect(83, 375, 91, 491)]
        header_rule = fitz.Rect(42, 45, 378, 63)
        page = FakePage(self.page_rect, chart + [header_rule])
        caption = [42.6, 577.9, 91.1, 587.1]
        result = _figure_crop_box(page, caption)
        self.assertIsNotNone(result)
        crop, orientation = result
        self.assertEqual(orientation, "above_caption")
        self.assertGreater(crop.y0, 300, "crop must exclude the header rule far above the chart")
        self.assertLess(crop.y1, 580, "crop must not extend below the caption")

    def test_full_page_frame_rect_is_ignored(self) -> None:
        frame = fitz.Rect(0, 0, self.page_rect.width, self.page_rect.height)
        chart = fitz.Rect(50, 300, 380, 500)
        page = FakePage(self.page_rect, [frame, chart])
        result = _figure_crop_box(page, [45.0, 250.0, 120.0, 262.0])
        self.assertIsNotNone(result)
        crop, _ = result
        self.assertGreater(crop.y0, 250)

    def test_clusters_split_disjoint_geometry(self) -> None:
        rects = [fitz.Rect(0, 0, 100, 100), fitz.Rect(0, 0, 80, 80), fitz.Rect(300, 300, 400, 400)]
        clusters = _cluster_rects(rects, gap=10.0)
        self.assertEqual(len(clusters), 2)

    def test_embedded_image_page_uses_fallback_and_trims_border(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            document = fitz.open()
            page = document.new_page(width=433, height=649)
            # A large dark-framed scan placed over the whole page.
            payload = Image.new("L", (300, 420), 255)
            for x in range(300):
                for y in (0, 1, 2, 417, 418, 419):
                    payload.putpixel((x, y), 0)
            for y in range(420):
                for x in (0, 1, 2, 297, 298, 299):
                    payload.putpixel((x, y), 0)
            import io

            buffer = io.BytesIO()
            payload.save(buffer, format="PNG")
            page.insert_image(page.rect, stream=buffer.getvalue())
            destination = root / "exhibit.png"
            from book_translation_pro.figures import _embedded_page_fallback

            metadata = _embedded_page_fallback(document, page, destination)
            self.assertIsNotNone(metadata)
            self.assertEqual(metadata["render_method"], "embedded_image_extract_trimmed")
            with Image.open(destination) as result:
                self.assertLessEqual(result.size, (300, 420))
            document.close()

    def test_trim_scan_border_crops_uniform_edges(self) -> None:
        image = Image.new("L", (200, 300), 30)
        inner = Image.new("L", (160, 240), 220)
        image.paste(inner, (20, 30))
        trimmed = _trim_scan_border(image)
        self.assertLess(trimmed.size[0], 200)
        self.assertLess(trimmed.size[1], 300)
        self.assertGreaterEqual(trimmed.size[0], 160)

    def test_inverted_lineart_detection_and_idempotence(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "map.png"
            Image.new("L", (60, 60), 10).save(path)
            self.assertTrue(looks_like_inverted_lineart(Image.open(path)))
            self.assertTrue(normalize_image_background(path))
            with Image.open(path) as flipped:
                self.assertEqual(flipped.getpixel((0, 0)), 245)
            # A normalized (bright) image is left untouched on rerun.
            self.assertFalse(normalize_image_background(path))


class PublicationLayoutTests(unittest.TestCase):
    def base_config(self) -> dict:
        return {
            "schema_version": 1,
            "project_id": "sample",
            "source": {"path": "source/original.pdf", "language": "en", "format": "pdf", "edition": ""},
            "target": {"language": "zh-CN"},
            "mode": "study",
            "outputs": ["pdf"],
            "workspace": ".",
        }

    def test_study_mode_defaults_to_symmetric_margins(self) -> None:
        options = _publication_options(self.base_config())
        self.assertFalse(options["mirror_margins"])
        self.assertEqual(options["gutter_mm"], 0)

    def test_publication_mode_keeps_mirror_margins(self) -> None:
        config = self.base_config()
        config["mode"] = "publication"
        options = _publication_options(config)
        self.assertTrue(options["mirror_margins"])
        self.assertEqual(options["gutter_mm"], 5)

    def test_explicit_mirror_override_wins(self) -> None:
        config = self.base_config()
        config["publishing"] = {"mirror_margins": True, "gutter_mm": 3}
        options = _publication_options(config)
        self.assertTrue(options["mirror_margins"])
        self.assertEqual(options["gutter_mm"], 3)

    def test_configure_settings_removes_mirror_when_disabled(self) -> None:
        from docx.oxml.ns import qn

        document = Document()
        _configure_settings(document, {"mirror_margins": True})
        self.assertIsNotNone(document.settings._element.find(qn("w:mirrorMargins")))
        _configure_settings(document, {"mirror_margins": False})
        self.assertIsNone(document.settings._element.find(qn("w:mirrorMargins")))
        self.assertIsNotNone(document.settings._element.find(qn("w:evenAndOddHeaders")))

    def test_table_widths_sum_to_text_measure(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=3)
        cells = (
            ("项目", "金额", "占收入比例"),
            ("材料成本", "$326,240", "34.2%"),
        )
        for row, values in zip(table.rows, cells):
            for cell, value in zip(row.cells, values):
                cell.text = value
        target = int((170 - 25 - 20) / 25.4 * 1440)
        widths = _table_column_widths(document, table, target)
        self.assertEqual(len(widths), 3)
        self.assertEqual(sum(widths), target)
        # The numeric and header columns must be wide enough for their text.
        for width in widths:
            self.assertGreaterEqual(width, 680)

    def test_detect_folio_position_prefers_top_when_folios_lead(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "source").mkdir()
            pdf_path = root / "source" / "original.pdf"
            document = fitz.open()
            for index in range(30):
                page = document.new_page(width=433, height=649)
                page.insert_text((40, 50), str(100 + index), fontsize=9)
                page.insert_text((40, 90), f"Body paragraph {index} " * 12, fontsize=10)
            document.save(pdf_path)
            document.close()
            config = self.base_config()
            self.assertEqual(detect_folio_position(root, config), "header")

    def test_detect_folio_position_footer_for_non_pdf(self) -> None:
        config = self.base_config()
        config["source"]["format"] = "epub"
        with tempfile.TemporaryDirectory() as temp:
            self.assertEqual(detect_folio_position(Path(temp), config), "footer")


class NumberConservationTests(unittest.TestCase):
    def test_english_magnitude_matches_cjk_magnitude(self) -> None:
        self.assertEqual(
            _number_tokens("a shortfall of 10 million workers"),
            _number_tokens("将短缺 1,000 万名工人"),
        )
        self.assertEqual(
            _number_tokens("revenue of $100,000"),
            _number_tokens("收入 10 万美元"),
        )

    def test_percent_and_decades_normalize(self) -> None:
        self.assertEqual(
            _number_tokens("15 percent of earnings"),
            _number_tokens("税前利润的 15%"),
        )
        self.assertEqual(
            _number_tokens("since the 1990s"),
            _number_tokens("自 20 世纪 90 年代以来"),
        )
        self.assertEqual(
            _number_tokens("in the early part of the twentieth century"),
            _number_tokens("20 世纪初"),
        )

    def test_inline_print_stamp_ignored(self) -> None:
        stamped = "fewer sunk costs if market ch03_4772.qxd 2/2/07 3:38 PM Page 55 continues"
        self.assertEqual(_number_tokens(stamped), _number_tokens("fewer sunk costs if market continues"))

    def test_real_numeric_change_still_detected(self) -> None:
        self.assertNotEqual(_number_tokens("lead time of 10 days"), _number_tokens("lead time of 11 days"))


class FigureTablePageBreakTests(unittest.TestCase):
    def _options(self) -> dict:
        return {"trim_width_mm": 170, "inside_margin_mm": 25, "outside_margin_mm": 20}

    def test_table_rows_cannot_split_and_keep_together(self) -> None:
        from docx.oxml.ns import qn

        document = Document()
        table = document.add_table(rows=3, cols=2)
        for row in table.rows:
            for cell in row.cells:
                cell.text = "x"
        _configure_tables(document, self._options())
        for index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            self.assertIsNotNone(tr_pr.find(qn("w:cantSplit")), f"row {index} missing cantSplit")
            keep = index < 2  # non-last rows keep with next
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self.assertEqual(paragraph.paragraph_format.keep_with_next, keep)
        self.assertIsNotNone(table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")))

    def test_figure_stays_with_caption(self) -> None:
        import io

        document = Document()
        buffer = io.BytesIO()
        Image.new("RGB", (8, 8), "white").save(buffer, format="PNG")
        buffer.seek(0)
        document.add_picture(buffer)
        document.add_paragraph("图表 6.8　海岸警卫队使命的重要性")
        document.add_paragraph("正文段落。")
        _keep_figures_with_captions(document)
        figure_paragraph = document.paragraphs[0]
        caption_paragraph = document.paragraphs[1]
        body_paragraph = document.paragraphs[2]
        self.assertTrue(figure_paragraph.paragraph_format.keep_with_next)
        self.assertTrue(figure_paragraph.paragraph_format.keep_together)
        self.assertTrue(caption_paragraph.paragraph_format.keep_with_next)
        self.assertFalse(body_paragraph.paragraph_format.keep_with_next)


if __name__ == "__main__":
    unittest.main()
