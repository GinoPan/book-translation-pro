from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from book_translation_pro.ingest import _promote_chm_linked_images
from book_translation_pro.publication import (
    _configure_tables,
    _format_title_page,
    _infer_target_title,
    _insert_original_title_line,
    _normalize_image_markup,
    _restore_promoted_image_assets,
    clean_markdown_for_publication,
    validate_docx_package,
    validate_epub_package,
    validate_publication_outputs,
)


class PublicationCleanupTests(unittest.TestCase):
    def test_valid_epub_topology_passes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "book.epub"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
                archive.writestr(
                    "META-INF/container.xml",
                    '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
                )
                archive.writestr(
                    "OEBPS/content.opf",
                    '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Test</dc:title><dc:language>zh-CN</dc:language></metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="ch1" href="ch1.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="ch1"/></spine></package>',
                )
                archive.writestr(
                    "OEBPS/nav.xhtml",
                    '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body><nav epub:type="toc"><ol><li><a href="ch1.xhtml#chapter-1">第一章</a></li></ol></nav></body></html>',
                )
                archive.writestr(
                    "OEBPS/ch1.xhtml",
                    '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1 id="chapter-1">正文</h1></body></html>',
                )
            result = validate_epub_package(path)
            self.assertTrue(result["passed"])

    def test_epub_missing_manifest_resource_is_blocking(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "book.epub"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
                archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="content.opf"/></rootfiles></container>')
                archive.writestr(
                    "content.opf",
                    '<package><metadata><title>Test</title><language>zh-CN</language></metadata><manifest><item id="nav" href="missing-nav.xhtml" properties="nav"/></manifest><spine/></package>',
                )
                archive.writestr("nav.xhtml", "<html/>")
            result = validate_epub_package(path)
            self.assertFalse(result["passed"])
            self.assertTrue(any("missing" in message for message in result["errors"]))

    def test_removes_print_markers_and_running_folios(self) -> None:
        source = """
<span>ch03_4772.qxd 2/2/07 3:38 PM Page 50</span>

**50** <span>**LEAN ACCOUNTING**</span>

正文中的 50% 必须保留。
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertNotIn("qxd", cleaned)
        self.assertNotIn("LEAN ACCOUNTING", cleaned)
        self.assertIn("50%", cleaned)

    def test_splits_hard_break_address_lines_and_removes_empty_fragments(self) -> None:
        source = "Publisher Street\\\nShanghai 200000\\\n\\\n***\n**\n"
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("Publisher Street\n\nShanghai 200000", cleaned)
        self.assertNotIn("***", cleaned)
        self.assertNotIn("\n\\\n", cleaned)

    def test_normalizes_html_image_references_for_all_outputs(self) -> None:
        source = r'<img src="assets\original/images/000000.jpg" class="calibre49" />'
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("![](assets/original/images/000000.jpg)", cleaned)
        self.assertNotIn("<img", cleaned)

    def test_restores_extracted_footnote_superscripts_without_changing_quantities(self) -> None:
        source = (
            '正文结尾<span class="calibre26">6</span> 后续文字。'
            '体现<span class="calibre26">11</span>个 *5S*。'
        )
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("正文结尾^6^ 后续文字。", cleaned)
        self.assertIn("体现 *5S*^11^。", cleaned)
        self.assertNotIn("11个 *5S*", cleaned)

    def test_joins_paragraph_split_after_dangling_conjunction(self) -> None:
        source = "这是一个跨页句子，因为\n\n后半句接在这里。\n\n下一段。"
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("这是一个跨页句子，因为后半句接在这里。", cleaned)
        self.assertNotIn("因为\n\n后半句", cleaned)
        self.assertIn("\n\n下一段。", cleaned)

    def test_replaces_static_toc_and_normalizes_headings(self) -> None:
        source = """
## **目录**

**第一章** 3

<span>**前言**</span>

前言正文。

# <span>**第一部分**</span>

**精益基础**

**1**

**第一章标题**

**1.1 第一节**

正文。
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("# 目录\n\n[[TOC]]", cleaned)
        self.assertNotIn("第一章** 3", cleaned)
        self.assertIn("# 前言", cleaned)
        self.assertIn("# 第一部分 精益基础", cleaned)
        self.assertIn("# 1 第一章标题", cleaned)
        self.assertIn("## 1.1 第一节", cleaned)

    def test_converts_raw_html_heading_for_docx_safe_markdown(self) -> None:
        cleaned = clean_markdown_for_publication(
            '<h1 style="page-break-before: always;"><span>Chapter 12</span></h1>'
        )
        self.assertEqual(cleaned, "# Chapter 12\n")

    def test_demotes_summary_prose_and_drops_recovery_tail(self) -> None:
        summary = "本段是被误识别为标题的内容，" * 12 + "应当作为正文。"
        source = f"""
# 第一部分 精益基础

# 1 第一章标题

## {summary}

## 1.1 第一节

正文。

# OCR 恢复的源页面

本单元仅包含源文件附带的 OCR 恢复页标题。

# 重复封底标题
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertIn(summary, cleaned)
        self.assertNotIn(f"## {summary}", cleaned)
        self.assertNotIn("OCR 恢复的源页面", cleaned)
        self.assertNotIn("重复封底标题", cleaned)

    def test_removes_editor_artifacts_and_splits_wide_tables(self) -> None:
        source = """
*** Add File: work/observations/unit-0001.json
{"schema_version":1,"new_entries":[]}

<!-- 内部视觉切片说明 -->

| 编号 | A | B | C | D | E | F |
|---:|---|---|---|---|---|---|
| 1 | a | b | c | d | e | f |
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertNotIn("Add File", cleaned)
        self.assertNotIn("schema_version", cleaned)
        self.assertNotIn("视觉切片", cleaned)
        self.assertIn("*（续表）*", cleaned)
        self.assertEqual(cleaned.count("| 编号 |"), 2)

    def test_keeps_compact_wide_measurement_table_together(self) -> None:
        source = """
| 总数 | 很差 | 差 | 中等 | 好 | 很好 | N/A |
|---:|---:|---:|---:|---:|---:|---:|
| WH #1（实施前） | 0 | 28 | 6 | 2 | 1 | 2 |
| WH #2（实施后） | 0 | 5 | 10 | 19 | 6 | 2 |
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertNotIn("续表", cleaned)
        self.assertEqual(cleaned.count("| 总数 |"), 1)

    def test_converts_html_wide_table_before_splitting(self) -> None:
        source = """
<table><thead><tr><th>A</th><th>B</th><th>C</th><th>D</th><th>E</th><th>F</th><th>G</th></tr></thead>
<tbody><tr><td>1</td><td>2</td><td>3</td><td>4</td><td>5</td><td>6</td><td>7</td></tr></tbody></table>
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertEqual(cleaned.count("| A |"), 2)
        self.assertIn("*（续表）*", cleaned)

    def test_splits_pandoc_short_alignment_separators(self) -> None:
        source = """
| A | B | C | D | E | F | G |
|:--|:--:|:--:|:--:|:--:|:--:|:--:|
| 1 | 2 | 3 | 4 | 5 | 6 | 7 |
"""
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("*（续表）*", cleaned)
        self.assertNotIn("| A | B | C | D | E | F | G |", cleaned)

    def test_reproportions_pandoc_short_separators(self) -> None:
        source = "| 中文表头 | Amount |\n|:--|:--:|\n| 一 | $2,043 |"
        cleaned = clean_markdown_for_publication(source)
        separator = cleaned.splitlines()[1]
        self.assertGreaterEqual(separator.count("-"), 160)

    def test_nested_toc_table_preserves_adjacent_information_table(self) -> None:
        source = (
            '<table><tr><td><table><tr><td><span id="TOC">目录</span></td></tr></table></td>'
            '<td><table><tr><th>Safari</th><th>说明</th></tr>'
            '<tr><td>浏览器</td><td>信息卡内容</td></tr></table></td></tr></table>'
        )
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("[[TOC]]", cleaned)
        self.assertIn("Safari", cleaned)
        self.assertIn("信息卡内容", cleaned)
        self.assertNotIn("| 目录 |", cleaned)
        self.assertLess(cleaned.index("[[TOC]]"), cleaned.index("Safari"))

    def test_drops_empty_html_table_and_extracts_prose_row(self) -> None:
        source = (
            "<table><tr><td>&nbsp;</td></tr></table>"
            "<table><tr><td>脚注说明：这是一条足够长的说明文字，应作为普通段落输出，不能撑宽表格列。</td>"
            "<td></td></tr></table>"
        )
        cleaned = clean_markdown_for_publication(source)
        self.assertIn("脚注说明", cleaned)
        self.assertNotIn("| 脚注说明", cleaned)

    def test_preserves_real_table_image_and_removes_platform_icon(self) -> None:
        source = (
            '<table><tr><td><img src="teamlib.gif"><img src="figures/real.png" alt="图"></td></tr></table>'
        )
        cleaned = clean_markdown_for_publication(source)
        self.assertNotIn("teamlib.gif", cleaned)
        self.assertIn("![图](assets/original/figures/real.png)", cleaned)

    def test_normalizes_bracketed_sup_marker_without_creating_note(self) -> None:
        cleaned = clean_markdown_for_publication("正文<sup>[38]</sup>继续")
        self.assertIn(r"^\[38\]^", cleaned)
        self.assertNotIn("^[38]^", cleaned)

    def test_title_inference_skips_front_matter_labels(self) -> None:
        source = "# 目录\n\n# 内容提要\n\n# 精益会计实务\n"
        self.assertEqual(_infer_target_title(source, "fallback", "zh-Hans"), "精益会计实务")

    def test_explicit_original_title_preserves_source_casing(self) -> None:
        source = "# 中文书名\n\n正文。\n"
        updated = _insert_original_title_line(source, "中文书名", "zh-Hans", "COS NFPA Handbook")
        self.assertIn("*COS NFPA Handbook*", updated)

    def test_title_page_has_a_hard_upper_bound_without_copyright_marker(self) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        document = Document()
        document.add_paragraph("中文书名", style="Title")
        document.add_heading("中文书名", level=1)
        body = [document.add_paragraph(f"正文段落 {index}") for index in range(45)]
        _format_title_page(document, {"title": "中文书名"})
        self.assertEqual(document.paragraphs[0].text, "中文书名")
        self.assertEqual(body[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertNotEqual(body[-1].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_table_configuration_uses_fixed_layout_without_rewriting_content(self) -> None:
        from docx import Document
        from docx.oxml.ns import qn

        document = Document()
        table = document.add_table(rows=2, cols=5)
        labels = ("编号", "COSO risk", "COSO", "COSO page", "IAD page")
        for cell, label in zip(table.rows[0].cells, labels):
            cell.text = label
        options = {"trim_width_mm": 170, "inside_margin_mm": 25, "outside_margin_mm": 20}
        _configure_tables(document, options)
        self.assertEqual(tuple(cell.text for cell in table.rows[0].cells), labels)
        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        self.assertIsNotNone(layout)
        self.assertEqual(layout.get(qn("w:type")), "fixed")
        children = list(table._tbl.tblPr)
        look = table._tbl.tblPr.find(qn("w:tblLook"))
        if look is not None:
            self.assertLess(children.index(layout), children.index(look))

    def test_publication_inventory_blocks_dropped_tables_and_images(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            markdown = root / "book.md"
            markdown.write_text(
                "This sentence is long enough to serve as a stable content anchor.\n\n"
                "| A | B |\n|---|---|\n| 1 | 2 |\n\n![](assets/original/figure.png)\n",
                encoding="utf-8",
            )
            docx = root / "book.docx"
            document = Document()
            document.add_paragraph("This sentence is long enough to serve as a stable content anchor.")
            document.save(docx)
            result = validate_publication_outputs(
                {"markdown": markdown, "docx": docx}, {"mode": "study"}
            )
            self.assertFalse(result["passed"])
            self.assertTrue(any("table inventory shrank" in error for error in result["errors"]))
            self.assertTrue(any("image inventory shrank" in error for error in result["errors"]))

    def test_docx_toc_is_required_only_when_the_master_requests_it(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            docx = Path(directory) / "book.docx"
            Document().save(docx)
            self.assertTrue(validate_docx_package(docx, require_toc=False)["passed"])
            result = validate_docx_package(docx, require_toc=True)
            self.assertFalse(result["passed"])
            self.assertIn("native Word TOC field is missing", result["errors"])


class ChmImagePromotionTests(unittest.TestCase):
    def test_promotes_linked_full_resolution_image(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "thumb.jpg").write_bytes(b"thumb")
            (root / "full.jpg").write_bytes(b"full-resolution-image")
            html = '<a href="full.jpg"><img src="thumb.jpg" width="350" height="200"></a>'
            promoted, count = _promote_chm_linked_images(html, root)
            self.assertEqual(count, 1)
            self.assertIn('src="full.jpg"', promoted)

    def test_normalizes_existing_markdown_thumbnail_link_to_full_image(self) -> None:
        source = '<a href="7485final/images/fig3-3_0.jpg">![点击展开](assets/original/7485final/images/fig3-3.jpg)</a>'
        cleaned = _normalize_image_markup(source)
        self.assertIn('assets/original/7485final/images/fig3-3_0.jpg', cleaned)
        self.assertNotIn('assets/original/7485final/images/fig3-3.jpg)', cleaned)

    def test_restores_missing_promoted_asset_from_chm_extract(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            source = project / "source-work" / "extracted" / "htmlz" / "fig.jpg"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"full-resolution")
            markdown = "![](assets/original/fig.jpg)"
            self.assertEqual(_restore_promoted_image_assets(project, markdown), 1)
            self.assertEqual((project / "assets" / "original" / "fig.jpg").read_bytes(), b"full-resolution")


if __name__ == "__main__":
    unittest.main()
